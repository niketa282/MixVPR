"""Two-page paper plan: findings, tables and figures (Word) plus LaTeX tables.

Run from the repo root with python-docx on PYTHONPATH:
    PYTHONPATH=~/pylib python results/doc_scripts/make_two_page_plan.py
"""
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SEEDS = [190223, 1, 2]
KEYS_GEN = ['pitts30k_test', 'pitts250k_test', 'msls_val', 'sf_xl_small', 'tokyo247', 'amstertime', 'st_lucia', 'eynsham']
KEYS_SEA = ['nordland', 'svox', 'svox_overcast', 'svox_rain', 'svox_snow', 'svox_sun']
NIGHT = 'svox_night'
ALL = KEYS_GEN + KEYS_SEA + [NIGHT]
GREY = RGBColor(0x5f, 0x5e, 0x58)


def load(suffix, seeds=SEEDS):
    d = pd.concat([pd.read_csv(f'results/seed{s}/test_results{suffix}.csv') for s in seeds])
    d = d[d.seed.isin(seeds)]
    p = d.pivot_table(index=['seed', 'condition'], columns='dataset', values='R1')
    g = pd.concat([p.loc[s] - p.loc[s].loc['none'] for s in seeds], keys=seeds)
    return p.groupby(level=1).mean(), g.groupby(level=1).mean()


(R_abs, R), (G_abs, G) = load(''), load('_gray')
groups = [('General', KEYS_GEN), ('Seasonal/weather', KEYS_SEA), ('Night', [NIGHT])]
ROWS = [('geo_k2', 'Mixed geometric (k = 2)'), ('phot_k3', 'Mixed photometric (k = 3)'), ('randaugment', 'RandAugment')]

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
doc.sections[0].left_margin = doc.sections[0].right_margin = Cm(2.0)
tno, fno = [0], [0]


def border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '000000')
    b.append(el)


def caption(text, kind='Table'):
    n = tno if kind == 'Table' else fno
    n[0] += 1
    p = doc.add_paragraph()
    r = p.add_run(f'{kind} {n[0]}. '); r.bold = True
    p.add_run(text)
    if kind == 'Table':
        p.paragraph_format.keep_with_next = True


def table(header, rows, widths, font=9.5, rules=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        border(c, 'top', 12); border(c, 'bottom', 6)
    for i, cells in enumerate(rows, 1):
        rc = t.rows[i].cells
        for j, text in enumerate(cells):
            par = rc[j].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            r = par.add_run(text); r.font.size = Pt(font)
        if cells[0] in rules:
            for c in rc: border(c, 'bottom', 4)
        if i == len(rows):
            for c in rc: border(c, 'bottom', 12)
    doc.add_paragraph()


def bullet(head, body):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(head + ' '); r.bold = True
    p.add_run(body)


def fig(path, width_cm, text):
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(text, kind='Figure')


f = lambda v: f'{v:+.1f}'
gm = lambda rep, c, ks: rep.loc[c, ks].mean()

# ------------------------------------------------------------------ intro
doc.add_heading('Two-page paper: recommended content', level=1)
doc.add_paragraph('Plan for a 2-page IEEE paper based on the completed MixVPR experiments: what to claim, which '
                  'figures and tables to use, and what to leave out. All numbers are Recall@1 (%), mean over three '
                  'seeds, from the 15-test-set evaluation.')
p = doc.add_paragraph()
r = p.add_run('Possible title: '); r.bold = True
p.add_run('"Colour, Not Darkness: How Augmentation and Grayscale Input Shape Night-Time Visual Place Recognition"')
p = doc.add_paragraph()
r = p.add_run('Contribution in one sentence: '); r.bold = True
p.add_run('In a controlled three-seed study of MixVPR on 15 benchmarks, geometric and photometric augmentation '
          'address different kinds of appearance change, their benefit depends on whether the input is RGB or '
          'grayscale, and the night-time failure of RGB models is driven by colour casts rather than darkness, which '
          'grayscale input largely removes.')

# ------------------------------------------------------------------ findings
doc.add_heading('Key findings to report', level=2)
bullet('1. Removing colour fixes most of the night problem.',
       f'SVOX Night R@1 rises from {R_abs.loc["none", NIGHT]:.1f} % (RGB) to {G_abs.loc["none", NIGHT]:.1f} % '
       f'(grayscale) without augmentation; the best model is grayscale with RandAugment '
       f'({G_abs.loc["randaugment", NIGHT]:.1f} %). Night queries are barely darker than overcast or snow queries, '
       'which RGB handles well, but they carry a strong orange colour cast (Table 2).')
bullet('2. The augmentation families address different changes (RGB).',
       f'Photometric augmentation helps at night (phot_k3 {f(R.loc["phot_k3", NIGHT])}, contrast '
       f'{f(R.loc["contrast", NIGHT])}), geometric augmentation hardly does (at most '
       f'{f(R.loc[["crop", "perspective", "rotate", "translate", "shear", "geo_k1", "geo_k2", "geo_k3"], NIGHT].max())}); '
       f'geometric helps more on urban viewpoint benchmarks (SF-XL: geo_k2 {f(R.loc["geo_k2", "sf_xl_small"])} vs. '
       f'phot_k3 {f(R.loc["phot_k3", "sf_xl_small"])}).')
bullet('3. The effect of augmentation depends on the input representation.',
       f'Every geometric augmentation helps grayscale more than RGB, but colour-related photometric gains shrink on '
       f'grayscale (contrast at night: {f(R.loc["contrast", NIGHT])} RGB vs. {f(G.loc["contrast", NIGHT])} grayscale) '
       f'while brightness only helps once colour is removed ({f(R.loc["brightness", NIGHT])} vs. '
       f'{f(G.loc["brightness", NIGHT])}).')
bullet('4. Combining transforms beats single transforms.',
       f'Mixed conditions give the most reliable gains: phot_k3 averages {f(R.loc["phot_k3", ALL].mean())} and geo_k2 '
       f'{f(R.loc["geo_k2", ALL].mean())} over 15 test sets (RGB), ahead of every single transform.')

# ------------------------------------------------------------------ table 1
doc.add_heading('Table for the paper: family × representation', level=2)
caption('R@1 gain (percentage points) over the no-augmentation model of the same input, mean over three seeds. '
        'General: 8 benchmarks; seasonal/weather: Nordland and 5 SVOX subsets; night: SVOX Night. First row: '
        'absolute R@1 of the no-augmentation model.')
rows = [['No augmentation (R@1)'] + [f'{R_abs.loc["none", ks].mean():.1f}' for _, ks in groups]
        + [f'{G_abs.loc["none", ks].mean():.1f}' for _, ks in groups]]
for c, lab in ROWS:
    rows.append([lab] + [f(gm(R, c, ks)) for _, ks in groups] + [f(gm(G, c, ks)) for _, ks in groups])
table(['', 'RGB gen.', 'RGB seas.', 'RGB night', 'Gray gen.', 'Gray seas.', 'Gray night'], rows,
      [4.6, 1.9, 1.9, 1.9, 1.9, 1.9, 1.9], font=9, rules=('No augmentation (R@1)', 'Mixed photometric (k = 3)'))
doc.add_paragraph('Fixed conditions (geo_k2, phot_k3) are used instead of the best condition per column, to avoid '
                  'picking winners per test-set group. The full per-condition results stay in the thesis.')

# ------------------------------------------------------------------ table 2 (colour statistics)
doc.add_heading('Table for the paper: why night is different', level=2)
cs = pd.read_csv('results/night_analysis/colour_stats_per_image.csv').groupby('set', sort=False).mean()
order = ['Database (day)', 'Overcast', 'Snow', 'Rain', 'Sun', 'Night']
rgb_night_r1 = {'Overcast': 'svox_overcast', 'Snow': 'svox_snow', 'Rain': 'svox_rain', 'Sun': 'svox_sun', 'Night': NIGHT}
caption('Colour statistics of SVOX images (400 per set, CIE L*a*b*): mean lightness L*, colour-cast strength '
        '√(ā*² + b̄*²) and yellow–blue component b̄*, with the R@1 of the RGB no-augmentation model.')
rows = []
for s in order:
    r1 = f'{R_abs.loc["none", rgb_night_r1[s]]:.1f}' if s in rgb_night_r1 else '—'
    rows.append([s, f'{cs.loc[s, "L"]:.0f}', f'{cs.loc[s, "cast"]:.1f}', f'{cs.loc[s, "b"]:+.1f}', r1])
table(['Set', 'Lightness L*', 'Cast strength', 'b* (yellow +)', 'RGB R@1'], rows, [3.6, 2.6, 2.8, 2.8, 2.4],
      font=9.5, rules=('Database (day)',))

# ------------------------------------------------------------------ figures
doc.add_heading('Figures for the paper', level=2)
fig('results/figures/night_rgb_vs_gray.png', 9.0,
    'SVOX Night R@1 for RGB and grayscale models: no augmentation, best geometric and best photometric condition '
    '(named), and RandAugment. Error bars: standard deviation over three seeds. One IEEE column '
    '(results/figures/night_rgb_vs_gray.pdf).')
fig('results/figures/vpr_pipeline_train.png', 16.5,
    'Training pipeline: augmentation (training only) and RGB or grayscale input. Full IEEE width '
    '(results/figures/vpr_pipeline_train.pdf). Optional if space is short.')

# ------------------------------------------------------------------ TrivialAugment
doc.add_heading('TrivialAugment: include only as a short side result', level=2)
d = pd.concat([pd.read_csv(f'results/seed{s}/test_results.csv') for s in SEEDS])
ta_seeds = sorted(set(d[d.condition == 'trivialaugment'].seed))
(_, R2) = load('', ta_seeds)
doc.add_paragraph(f'TrivialAugment has {len(ta_seeds)} completed seed(s) ({", ".join(map(str, ta_seeds))}); the third '
                  'run failed and is being repeated. It is not part of the research question (it is, like RandAugment, '
                  'a mixed automatic policy), so report it in one or two sentences, marked as preliminary if the '
                  'third seed is not finished.')
caption(f'RandAugment vs. TrivialAugment (RGB), R@1 gain over no augmentation, mean over seeds '
        f'{", ".join(map(str, ta_seeds))}.')
rows = []
for lab, ks in [('All 15 test sets', ALL), ('Nordland', ['nordland']), ('SVOX Night', [NIGHT])]:
    rows.append([lab, f(R2.loc['randaugment', ks].mean()), f(R2.loc['trivialaugment', ks].mean())])
table(['', 'RandAugment', 'TrivialAugment'], rows, [4.5, 3.2, 3.2], font=9.5)
doc.add_paragraph(f'Suggested sentence: "TrivialAugment, a tuning-free alternative, trained about 20 % faster but '
                  f'gave smaller gains than RandAugment ({f(R2.loc["trivialaugment", ALL].mean())} vs. '
                  f'{f(R2.loc["randaugment", ALL].mean())} R@1 on average) and degraded seasonal recognition '
                  f'(Nordland {f(R2.loc["trivialaugment", "nordland"])})."')

# ------------------------------------------------------------------ layout + leave out
doc.add_heading('Suggested layout (2 pages, IEEE)', level=2)
for t in ['Page 1: introduction and research question (short), method in one paragraph (MixVPR, GSV-Cities, 50 '
          'epochs, 3 seeds, augmentations listed, RGB vs. grayscale, 15 test sets, gain vs. no augmentation), '
          'optional pipeline figure (full width, top).',
          'Page 2: Table 1 (family × representation), Figure: night RGB vs. grayscale, Table 2 (colour statistics) '
          'with findings 1–4, one sentence on TrivialAugment, short conclusion.',
          'If space is short: drop the pipeline figure first, then merge Table 2 into the text.']:
    doc.add_paragraph(t, style='List Bullet')
doc.add_heading('Leave out of the 2 pages', level=2)
for t in ['Per-test-set tables and full heatmaps (keep for the thesis).',
          'The strength sweep and the night test-time experiment (not finished).',
          'Validation curves and the 30 vs. 50 epoch discussion (one clause in the method is enough).',
          'Single-seed or 2-seed results presented as final.']:
    doc.add_paragraph(t, style='List Bullet')

doc.save('results/two_page_paper_plan.docx')

# ------------------------------------------------------------------ LaTeX tables
tex = ['% requires \\usepackage{booktabs} and \\usepackage{graphicx}', '',
       '\\begin{table}[t]', '\\centering',
       '\\caption{R@1 gain over no augmentation (same input), mean of 3 seeds. Gen.: 8 general benchmarks; '
       'Seas.: Nordland + 5 SVOX subsets; Night: SVOX Night. First row: absolute R@1.}',
       '\\label{tab:family_repr}', '\\setlength{\\tabcolsep}{3pt}', '\\resizebox{\\columnwidth}{!}{%',
       '\\begin{tabular}{@{}lrrrrrr@{}}', '\\toprule',
       ' & \\multicolumn{3}{c}{RGB} & \\multicolumn{3}{c}{Grayscale} \\\\',
       '\\cmidrule(lr){2-4} \\cmidrule(l){5-7}', ' & Gen. & Seas. & Night & Gen. & Seas. & Night \\\\', '\\midrule']
cells = [f'{R_abs.loc["none", ks].mean():.1f}' for _, ks in groups] + [f'{G_abs.loc["none", ks].mean():.1f}' for _, ks in groups]
tex.append('No aug.\\ (R@1) & ' + ' & '.join(cells) + ' \\\\')
tex.append('\\midrule')
for c, lab in ROWS:
    lab = lab.replace('(k = 2)', '($k{=}2$)').replace('(k = 3)', '($k{=}3$)')
    vals = [gm(R, c, ks) for _, ks in groups] + [gm(G, c, ks) for _, ks in groups]
    tex.append(f'{lab} & ' + ' & '.join(f(v).replace('-', '$-$') for v in vals) + ' \\\\')
    if c == 'phot_k3':
        tex.append('\\midrule')
tex += ['\\bottomrule', '\\end{tabular}%', '}', '\\end{table}', '',
        '\\begin{table}[t]', '\\centering',
        '\\caption{Colour statistics of SVOX images (400 per set, CIE L*a*b*) and R@1 of the RGB model without '
        'augmentation. Night is not darker than overcast or snow but has a strong yellow/orange cast.}',
        '\\label{tab:colour}', '\\setlength{\\tabcolsep}{4pt}', '\\begin{tabular}{@{}lrrrr@{}}', '\\toprule',
        'Set & $L^*$ & Cast & $b^*$ & R@1 \\\\', '\\midrule']
for s in order:
    r1 = f'{R_abs.loc["none", rgb_night_r1[s]]:.1f}' if s in rgb_night_r1 else '--'
    b = f'{cs.loc[s, "b"]:+.1f}'.replace('-', '$-$')
    tex.append(f'{s} & {cs.loc[s, "L"]:.0f} & {cs.loc[s, "cast"]:.1f} & {b} & {r1} \\\\')
    if s == 'Database (day)':
        tex.append('\\midrule')
tex += ['\\bottomrule', '\\end{tabular}', '\\end{table}', '']
open('results/two_page_tables.tex', 'w').write('\n'.join(tex))
print('saved results/two_page_paper_plan.docx and results/two_page_tables.tex')
