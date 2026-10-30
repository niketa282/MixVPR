"""RGB strength sweep document: table, figure and findings (results/strength_sweep_rgb.docx).

Run from the repo root with python-docx on PYTHONPATH, after
    python results/plot_strength_sweep.py --full
"""
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SEEDS = [190223, 1, 2]
GEN = ['pitts30k_test', 'pitts250k_test', 'msls_val', 'sf_xl_small', 'tokyo247', 'amstertime', 'st_lucia', 'eynsham']
SEA = ['nordland', 'svox', 'svox_overcast', 'svox_rain', 'svox_snow', 'svox_sun']
ALL = GEN + SEA + ['svox_night']
ROWS = [('Perspective 0.1', 'perspective_0.1'), ('Perspective 0.2 (default)', 'perspective'),
        ('Perspective 0.3', 'perspective_0.3'), ('Perspective 0.4', 'perspective_0.4'),
        ('Contrast ±0.2', 'contrast_0.8-1.2'), ('Contrast ±0.4 (default)', 'contrast'),
        ('Contrast ±0.6', 'contrast_0.4-1.6'), ('Contrast ±0.8', 'contrast_0.2-1.8')]

d = pd.concat([pd.read_csv(f'results/seed{s}/test_results.csv') for s in SEEDS])
p = d.pivot_table(index=['seed', 'condition'], columns='dataset', values='R1')
g = pd.concat([p.loc[s] - p.loc[s].loc['none'] for s in SEEDS], keys=SEEDS)
gm = g.groupby(level=1).mean()
f = lambda v: '0.0' if round(v, 2) == 0 else f'{v:+.2f}'
f1 = lambda v: '0.0' if round(v, 1) == 0 else f'{v:+.1f}'
G = lambda c, ks: gm.loc[c, ks].mean()
SD = lambda c, ks: g.xs(c, level=1)[ks].mean(axis=1).std()
n_up = lambda c: int((gm.loc[c, ALL] > 0).sum())
s1 = lambda c, k: f1(gm.loc[c, k])

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
doc.sections[0].left_margin = doc.sections[0].right_margin = Cm(2.0)


def border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '000000')
    b.append(el)


def caption(text, kind):
    p_ = doc.add_paragraph()
    r = p_.add_run(f'{kind}. '); r.bold = True
    p_.add_run(text)
    p_.paragraph_format.keep_with_next = kind.startswith('Table')


def table(header, rows, widths, bold_rows=(), rule_after=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        border(c, 'top', 12); border(c, 'bottom', 6)
    for i, cells in enumerate(rows, 1):
        rc = t.rows[i].cells
        for j, text in enumerate(cells):
            par = rc[j].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.keep_with_next = i < len(rows)
            r = par.add_run(text); r.font.size = Pt(9.5); r.bold = (i - 1) in bold_rows
        if i in rule_after:
            for c in rc: border(c, 'bottom', 4)
        if i == len(rows):
            for c in rc: border(c, 'bottom', 12)


def bullet(parts):
    p_ = doc.add_paragraph(style='List Bullet')
    for text, bold in parts:
        r = p_.add_run(text); r.bold = bold


# ------------------------------------------------------------------ title + setup
doc.add_heading('Strength sweep (RGB): perspective and contrast', level=1)
doc.add_paragraph(
    'Perspective and contrast are the single geometric and photometric operations with the highest average gain on RGB '
    'input, so their strength was varied around the default used in the main experiments (perspective distortion '
    'scale 0.2; contrast factor 0.6–1.4, i.e. ±0.4 around 1). Each condition was trained on GSV-Cities with RGB input '
    'for 50 epochs with three seeds and evaluated on all 15 test sets. Values are the change in Recall@1 (percentage '
    'points) from the no-augmentation RGB model, computed per seed and averaged; "All 15" also gives the standard '
    'deviation over seeds. "Sets improved" counts test sets with a positive mean gain.')

# ------------------------------------------------------------------ purpose
P, P3, P4 = 'perspective', 'perspective_0.3', 'perspective_0.4'
C, C2, C6, C8 = 'contrast', 'contrast_0.8-1.2', 'contrast_0.4-1.6', 'contrast_0.2-1.8'
doc.add_heading('Purpose of the experiment', level=2)
doc.add_paragraph(
    'In the main experiments every operation used one fixed, fairly mild strength. The sweep checks whether the main '
    'conclusions depend on these chosen strengths, i.e. whether geometric vs. photometric augmentation would compare '
    'differently with other settings. It addresses three questions:')
bullet([('Were the fixed settings reasonable? ', True),
        ('If a default strength were far from its best value, the main comparison between families could be '
         'misleading.', False)])
bullet([('How does strength itself affect robustness? ', True),
        ('Strength is a property of an augmentation operation, so how gains change with strength is part of how '
         'operations affect VPR robustness.', False)])
bullet([("Is RandAugment's lead explained by strength? ", True),
        ('If stronger single operations reached RandAugment, its advantage would come from strength; if not, it is '
         'more likely to come from the diversity of operations it combines.', False)])

# ------------------------------------------------------------------ table
caption('R@1 gain over no augmentation for each strength (RGB, mean over 3 seeds). Bold: the best strength of '
        'each operation over all 15 test sets.', 'Table 1')
rows = []
for lab, c in ROWS:
    rows.append([lab, f'{f(G(c, ALL))} ± {SD(c, ALL):.2f}', f(G(c, GEN)), f(G(c, SEA)), f1(gm.loc[c, 'nordland']),
                 f'{f1(gm.loc[c, "svox_night"])} ± {SD(c, ["svox_night"]):.1f}', f'{n_up(c)}/15'])
best_p = max(range(4), key=lambda i: G(ROWS[i][1], ALL))
best_c = 4 + max(range(4), key=lambda i: G(ROWS[4 + i][1], ALL))
table(['Condition', 'All 15', 'General (8)', 'Seas./weather (6)', 'Nordland', 'Night', 'Sets improved'], rows,
      [4.0, 2.4, 1.9, 2.3, 1.7, 2.2, 1.9], bold_rows=(best_p, best_c), rule_after=(4,))
doc.add_paragraph()

# ------------------------------------------------------------------ figure
doc.add_picture('results/figures/strength_sweep_rgb_full.png', width=Cm(16.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption('R@1 gain over no augmentation against augmentation strength, for each column of Table 1 (mean over 3 '
        'seeds; shaded band = ± one standard deviation over seeds; ring = default strength of the main experiments).',
        'Figure 1')

# ------------------------------------------------------------------ findings
doc.add_heading('Findings', level=2)

doc.add_heading('1. Perspective has a sweet spot, and the default (0.2) is it', level=3)
bullet([('Gains rise from 0.1 to 0.2 and then fall. ', False),
        (f'At 0.4 the average gain is gone ({f(G(P4, ALL))}, only {n_up(P4)}/15 sets improved)', True),
        ('; too much distortion produces unrealistic training images.', False)])
bullet([('There is a trade-off between test sets: ', False),
        ('stronger perspective keeps helping the viewpoint-heavy sets', True),
        (f' (SF-XL {s1(P4, "sf_xl_small")}, MSLS {s1(P4, "msls_val")} at 0.4) but hurts Tokyo 24/7 '
         f'({s1(P4, "tokyo247")}) and the weather sets (rain {s1(P4, "svox_rain")}, snow {s1(P4, "svox_snow")}).',
         False)])
bullet([('At night perspective has no reliable effect at any strength', True),
        (f': the spread between seeds (about ±{SD(P, ["svox_night"]):.0f} points at the default) is larger than '
         f'the differences between strengths.', False)])

doc.add_heading('2. Contrast: stronger is better, and the default was too weak', level=3)
bullet([('Night gains rise steadily with strength ', False),
        (f'({f1(G(C2, ["svox_night"]))} → {f1(G(C, ["svox_night"]))} → {f1(G(C6, ["svox_night"]))} → '
         f'{f1(G(C8, ["svox_night"]))}).', True)])
bullet([('The strongest setting (±0.8) doubles the overall gain', True),
        (f' ({f(G(C8, ALL))} vs. {f(G(C, ALL))} at the default) and helps across the board: Tokyo 24/7 '
         f'{s1(C8, "tokyo247")}, SVOX Sun {s1(C8, "svox_sun")}, SVOX Rain {s1(C8, "svox_rain")}.', False)])
bullet([('A single contrast operation at ±0.8 nearly matches mixed photometric augmentation', True),
        (f' (k = 3: {f(G("phot_k3", ALL))} overall, {f1(G("phot_k3", ["svox_night"]))} at night), but remains well '
         f'below RandAugment ({f(G("randaugment", ALL))} overall, {f1(G("randaugment", ["svox_night"]))} at night).',
         False)])

doc.add_heading('3. One noisy point', level=3)
bullet([(f'Contrast ±0.6 has a large spread between seeds (±{SD(C6, ALL):.2f} over all 15 sets), so its dip below '
         f'the trend probably comes from one seed rather than a real effect.', False)])

# ------------------------------------------------------------------ comparison with RandAugment
ORIG = ['crop', 'perspective', 'rotate', 'translate', 'shear', 'geo_k1', 'geo_k2', 'geo_k3', 'brightness', 'contrast',
        'autocontrast', 'blur', 'phot_k1', 'phot_k2', 'phot_k3']
SWEEP = [c for _, c in ROWS if c not in ORIG]
gr = gm.round(1)  # compare at the precision the results are reported
beats = lambda conds: [k for k in ALL if gr.loc[conds, k].max() > gr.loc['randaugment', k]]
NAMES = {'sf_xl_small': 'SF-XL', 'eynsham': 'Eynsham', 'nordland': 'Nordland', 'msls_val': 'MSLS',
         'tokyo247': 'Tokyo 24/7', 'amstertime': 'AmsterTime', 'st_lucia': 'St Lucia', 'pitts30k_test': 'Pitts30k',
         'pitts250k_test': 'Pitts250k', 'svox': 'SVOX', 'svox_overcast': 'SVOX Overcast', 'svox_rain': 'SVOX Rain',
         'svox_snow': 'SVOX Snow', 'svox_sun': 'SVOX Sun', 'svox_night': 'SVOX Night'}
b_orig, b_all, b_c8 = beats(ORIG), beats(ORIG + SWEEP), beats([C8])

doc.add_heading('Comparison with RandAugment', level=2)
ra_all, ra_n = G('randaugment', ALL), G('randaugment', ['svox_night'])
caption('Best single-operation and mixed conditions compared with RandAugment (RGB, mean over 3 seeds).', 'Table 2')
cmp_rows = [('RandAugment', 'randaugment'), ('Mixed photometric, k = 3 (best original condition)', 'phot_k3'),
            ('Contrast ±0.8 (best sweep condition)', C8), ('Contrast ±0.4 (default)', C)]
table(['Condition', 'All 15', 'Night', "Share of RandAugment's gain"],
      [[lab, f(G(c, ALL)), f1(G(c, ['svox_night'])), f'{100 * G(c, ALL) / ra_all:.0f}%'] for lab, c in cmp_rows],
      [7.6, 2.4, 2.2, 4.2])
doc.add_paragraph()
bullet([('Stronger augmentation narrowed the gap but did not close it. ', True),
        (f'Contrast at ±0.8 doubles what a single operation achieves ({100 * G(C, ALL) / ra_all:.0f}% → '
         f'{100 * G(C8, ALL) / ra_all:.0f}% of RandAugment\'s gain) but remains below half of it, and the gap at night '
         f'stays large ({f1(G(C8, ["svox_night"]))} vs. {f1(ra_n)}).', False)])
bullet([('No new test sets where RandAugment is beaten. ', True),
        (f'Before the sweep, a geometric or photometric condition beat RandAugment on {len(b_orig)}/15 RGB test sets '
         f'({", ".join(NAMES[k] for k in b_orig)}); including the sweep conditions it is still {len(b_all)}/15. '
         f'Contrast ±0.8 on its own beats RandAugment on {len(b_c8)}/15 sets.', False)])
bullet([('Interpretation: ', True),
        ("RandAugment's advantage, especially at night, does not come from strength alone. A single operation at "
         'maximum strength covers one kind of change, whereas RandAugment combines three operations per image drawn '
         'from about 14 geometric and colour operations. This is consistent with the finding that mixing operations '
         'helps more than single operations, and with the night images needing colour variation rather than only '
         'contrast.', False)])

doc.add_heading('Takeaway', level=2)
doc.add_paragraph(
    'The two families respond differently to strength. Geometric augmentation has an optimum: moderate perspective '
    'helps, strong perspective helps viewpoint-heavy test sets but hurts the weather sets. Photometric '
    'augmentation for RGB keeps improving within the tested range, consistent with the night finding that RGB models '
    'benefit from as much lighting and contrast variation as possible.')
doc.add_paragraph(
    'Answering the three questions: the fixed settings were reasonable for perspective (the default is the best '
    'strength tested) but too weak for contrast, so the main experiments, if anything, underestimate photometric '
    'augmentation for RGB; the main conclusions do not change. The two families respond differently to strength. '
    "RandAugment's lead is not explained by strength, pointing to the diversity of its operations.")

doc.add_heading('Suggested wording for the thesis', level=2)
p_ = doc.add_paragraph(
    'To verify that our conclusions do not depend on the chosen augmentation strengths, we vary the strength of the '
    'best single geometric (perspective) and photometric (contrast) operation on RGB input. Perspective peaks at the '
    'default strength, whereas contrast continues to improve, especially at night '
    f'({f1(G(C, ["svox_night"]))} → {f1(G(C8, ["svox_night"]))} R@1). Increasing strength roughly doubles the gain of '
    f'single-operation contrast augmentation ({f(G(C, ALL))} → {f(G(C8, ALL))} R@1 over 15 test sets), but it '
    f'remains below half of RandAugment\'s gain ({f(ra_all)}) and does not outperform RandAugment on any additional '
    "test set, indicating that RandAugment's advantage stems from the diversity of its operations rather than their "
    'strength.')
p_.runs[0].italic = True

out = 'results/strength_sweep_rgb.docx'
doc.save(out)
print('saved', out)
