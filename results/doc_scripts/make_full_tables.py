"""RGB + grayscale full tables, gain tables, interaction table and key findings. Run from MixVPR."""
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

S = [190223, 1, 2]
SINGLE = ['crop', 'perspective', 'rotate', 'translate', 'shear']
MIXED = ['geo_k1', 'geo_k2', 'geo_k3']
AUGS = SINGLE + MIXED
CONDS = ['none'] + AUGS
LABEL = {'none': 'No augmentation', 'crop': 'Random resized crop', 'perspective': 'Perspective transform',
         'rotate': 'Rotation', 'translate': 'Translation', 'shear': 'Shear',
         'geo_k1': 'Mixed, k = 1', 'geo_k2': 'Mixed, k = 2', 'geo_k3': 'Mixed, k = 3'}
GEN = [('pitts30k_test', 'Pitts30k'), ('pitts250k_test', 'Pitts250k'), ('msls_val', 'MSLS'),
       ('sf_xl_small', 'SF-XL'), ('tokyo247', 'Tokyo 24/7'), ('amstertime', 'AmsterTime'),
       ('st_lucia', 'St Lucia'), ('eynsham', 'Eynsham')]
SEA = [('nordland', 'Nordland'), ('svox', 'SVOX (all)'), ('svox_night', 'SVOX Night'),
       ('svox_overcast', 'SVOX Overcast'), ('svox_rain', 'SVOX Rain'), ('svox_snow', 'SVOX Snow'),
       ('svox_sun', 'SVOX Sun')]
ALL = GEN + SEA
K_GEN, K_SEA, K_ALL = [k for k, _ in GEN], [k for k, _ in SEA], [k for k, _ in ALL]
NAME = dict(ALL)
GREY = RGBColor(0x5f, 0x5e, 0x58)


def load(suffix):
    d = pd.concat([pd.read_csv(f'results/seed{s}/test_results{suffix}.csv') for s in S])
    d = d[d.condition.isin(CONDS)]
    out = {}
    for m in ('R1', 'R5', 'R10'):
        p = d.pivot_table(index=['seed', 'condition'], columns='dataset', values=m)[K_ALL]
        g = pd.concat([p.loc[s] - p.loc[s].loc['none'] for s in S], keys=S)
        out[m] = p.groupby(level=1).mean().loc[CONDS]
        out['d' + m] = g.groupby(level=1).mean().loc[CONDS]
        if m == 'R1':
            out['runs'] = p
            out['pos'] = g.groupby(level=1).agg(lambda v: int((v > 0).sum())).loc[CONDS]
            out['up'] = g.groupby(level=1).agg(lambda v: bool((v > 0).all())).loc[CONDS]
            out['dn'] = g.groupby(level=1).agg(lambda v: bool((v < 0).all())).loc[CONDS]
            out['g_runs'] = g
    return out


rgb, gray = load(''), load('_gray')

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
doc.sections[0].left_margin = doc.sections[0].right_margin = Cm(2.0)
tno = [0]


def border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '000000')
    b.append(el)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)


def caption(text):
    tno[0] += 1
    p = doc.add_paragraph()
    r = p.add_run(f'Table {tno[0]}. '); r.bold = True
    p.add_run(text)
    p.paragraph_format.keep_with_next = True


def table(header, rows, widths, font=9, rule_rows=(), bold=(), group_rows=(), align_left=1):
    """rows: list of lists of cells (str or list of (text, small)). group_rows: row indexes that are block titles."""
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    repeat_header(t.rows[0])
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j < align_left else WD_ALIGN_PARAGRAPH.CENTER
        border(c, 'top', 12); border(c, 'bottom', 6)
    for i, cells in enumerate(rows, 1):
        rc = t.rows[i].cells
        if (i - 1) in group_rows:
            merged = rc[0].merge(rc[-1])
            r = merged.paragraphs[0].add_run(cells[0]); r.bold = True; r.font.size = Pt(font)
            merged.paragraphs[0].paragraph_format.space_before = Pt(4)
            merged.paragraphs[0].paragraph_format.keep_with_next = True
            continue
        for j, cell in enumerate(cells):
            parts = cell if isinstance(cell, list) else [(cell, False)]
            par = rc[j].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j < align_left else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            for text, small in parts:
                r = par.add_run(text); r.font.size = Pt(font - 1.5 if small else font)
                r.bold = (i, j) in bold and not small
                if small:
                    r.font.color.rgb = GREY
        if (i - 1) in rule_rows:
            for c in rc: border(c, 'bottom', 4)
        if i == len(rows):
            for c in rc: border(c, 'bottom', 12)
    doc.add_paragraph()


def section(landscape):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    w, h = s.page_width, s.page_height
    if (landscape and w < h) or (not landscape and w > h):
        s.page_width, s.page_height = h, w
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.left_margin = s.right_margin = Cm(1.8 if landscape else 2.0)


def f(v, sign=False):
    if sign:
        return '0.00' if round(v, 2) == 0 else f'{v:+.2f}'
    return f'{v:.2f}'


# ---------------------------------------------------------------- intro
doc.add_heading('MixVPR geometric augmentation: full results for RGB and grayscale', level=1)
doc.add_paragraph('MixVPR (ResNet-50), trained on GSV-Cities for 50 epochs, three seeds per condition (190223, 1, 2); '
                  'checkpoint chosen on Pitts30k-val Recall@1. RGB models use colour images throughout; grayscale '
                  'models use grayscale images (three identical channels) in training, validation and testing.')
for t in [
    'R@K: Recall@K (%), mean over the three seeds.',
    'ΔR@K: change from the no-augmentation model of the same representation (RGB vs. RGB, grayscale vs. '
    'grayscale), computed within each seed and then averaged. Positive = better.',
    'Mixed, k = 1/2/3: k of the five single transforms applied to each training image, chosen at random per image.',
    'Bold in Tables 1–4: highest R@1 for that test set. Bold in Tables 5–8: same direction in all three seeds.',
]:
    doc.add_paragraph(t, style='List Bullet')

# ---------------------------------------------------------------- Tables 1-4: full per-test-set tables
for rep, name in [(rgb, 'RGB'), (gray, 'Grayscale')]:
    for group, gname in [(GEN, 'general benchmarks'), (SEA, 'seasonal and weather/lighting change')]:
        section(False)
        doc.add_heading(f'{name}: {gname}', level=2)
        caption(f'{name} results on the {gname}. Recall@K (%) and change from no augmentation (ΔR@K, percentage '
                'points); mean over three seeds.')
        rows, group_rows, rule_rows, bold = [], set(), set(), set()
        for k, label in group:
            group_rows.add(len(rows))
            nd = None
            rows.append([label])
            best = rep['R1'][k].max()
            for c in CONDS:
                rows.append([LABEL[c], f(rep['R1'].loc[c, k]), f(rep['R5'].loc[c, k]), f(rep['R10'].loc[c, k]),
                             f(rep['dR1'].loc[c, k], True), f(rep['dR5'].loc[c, k], True)])
                if round(rep['R1'].loc[c, k], 2) == round(best, 2):
                    bold.add((len(rows), 1))
                if c in ('none', 'shear'):
                    rule_rows.add(len(rows) - 1)
        table(['Training condition', 'R@1', 'R@5', 'R@10', 'ΔR@1', 'ΔR@5'], rows,
              [5.0, 2.1, 2.1, 2.1, 2.1, 2.1], font=9, rule_rows=rule_rows, bold=bold, group_rows=group_rows)

# ---------------------------------------------------------------- Tables 5-8: gains (workbook style)
for rep, name in [(rgb, 'RGB'), (gray, 'Grayscale')]:
    for group, gname in [(GEN, 'general benchmarks'), (SEA, 'seasonal and weather/lighting change')]:
        section(True)
        doc.add_heading(f'{name}: change from no augmentation on the {gname}', level=2)
        caption(f'{name}: change in R@1 (percentage points) from no augmentation on the {gname}, mean over three '
                'seeds; brackets: seeds with a positive change. Bold: same direction in all three seeds.')
        keys = [k for k, _ in group]
        rows, bold = [], set()
        for i, c in enumerate(AUGS, 1):
            rows.append([LABEL[c]] + [[(f(rep['dR1'].loc[c, k], True), False), (f' ({rep["pos"].loc[c, k]}/3)', True)]
                                      for k in keys])
            for j, k in enumerate(keys, 1):
                if rep['up'].loc[c, k] or rep['dn'].loc[c, k]:
                    bold.add((i, j))
        table(['Condition'] + [s for _, s in group], rows, [4.0] + [(25.6 - 4.0) / len(keys)] * len(keys),
              font=8.5, rule_rows={len(SINGLE) - 1}, bold=bold)

# ---------------------------------------------------------------- interaction tables
section(False)
doc.add_heading('Interaction between augmentation and input representation', level=2)
doc.add_paragraph('ΔR@1 RGB and ΔR@1 Gray are each augmentation\'s average R@1 gain over its own no-augmentation '
                  'baseline. The interaction difference is ΔR@1 Gray − ΔR@1 RGB: positive means the augmentation '
                  'helps grayscale models more than RGB models. "Sets" counts test sets (of 15) where the grayscale '
                  'gain is larger than the RGB gain.')


def inter(keys):
    out = {}
    for c in AUGS:
        r, g = rgb['dR1'].loc[c, keys], gray['dR1'].loc[c, keys]
        out[c] = (r.mean(), g.mean(), g.mean() - r.mean(), int((g > r).sum()))
    return out


I_all = inter(K_ALL)
caption('Augmentation × representation interaction, averaged over all 15 test sets (R@1 gain, percentage points).')
rows = [[LABEL[c], f(v[0], True), f(v[1], True), f(v[2], True), f'{v[3]} of 15'] for c, v in I_all.items()]
table(['Augmentation', 'ΔR@1 RGB', 'ΔR@1 Gray', 'Interaction difference', 'Sets gray > RGB'], rows,
      [4.6, 2.5, 2.5, 3.6, 3.0], font=9.5, rule_rows={len(SINGLE) - 1})

I_gen, I_sea = inter(K_GEN), inter(K_SEA)
I_nonight = inter([k for k in K_ALL if k != 'svox_night'])
caption('Interaction difference split by test-set group, and over all test sets without SVOX Night (the test set '
        'with by far the largest RGB–grayscale gap).')
rows = []
for c in AUGS:
    rows.append([LABEL[c],
                 f(I_gen[c][0], True), f(I_gen[c][1], True), f(I_gen[c][2], True),
                 f(I_sea[c][0], True), f(I_sea[c][1], True), f(I_sea[c][2], True),
                 f(I_nonight[c][2], True)])
table(['Augmentation', 'Gen. RGB', 'Gen. Gray', 'Gen. diff.', 'Sea. RGB', 'Sea. Gray', 'Sea. diff.', 'Diff. w/o Night'],
      rows, [3.6, 1.7, 1.7, 1.7, 1.7, 1.7, 1.7, 2.2], font=8.5, rule_rows={len(SINGLE) - 1})

# ---------------------------------------------------------------- key findings
doc.add_heading('Key findings and general patterns', level=2)
none_gap = gray['R1'].loc['none'] - rgb['R1'].loc['none']
up_r = rgb['up'].loc[AUGS].sum(axis=1); up_g = gray['up'].loc[AUGS].sum(axis=1)
dn_r = rgb['dn'].loc[AUGS].sum(axis=1); dn_g = gray['dn'].loc[AUGS].sum(axis=1)
avg_r = rgb['dR1'].loc[AUGS].mean(axis=1); avg_g = gray['dR1'].loc[AUGS].mean(axis=1)
night_aug_gap = (gray['R1'].loc[AUGS, 'svox_night'] - rgb['R1'].loc[AUGS, 'svox_night'])
inter_vals = pd.Series({c: v[2] for c, v in I_all.items()})
findings = [
    ('Geometric augmentation helps, modestly and broadly.',
     f'On RGB, every augmentation except crop improves R@1 on average (from {avg_r.drop("crop").min():+.2f} to '
     f'{avg_r.max():+.2f} points over 15 test sets). Gains are largest on SF-XL, MSLS, Nordland and SVOX Sun, small '
     'but consistent on the Pittsburgh sets, and absent on near-ceiling sets (St Lucia, SVOX all, SVOX Overcast).'),
    ('Combining transforms is the most reliable option.',
     f'Mixed k = 2 has the largest average RGB gain ({avg_r["geo_k2"]:+.2f}) and is consistently better on '
     f'{up_r["geo_k2"]} of 15 test sets and never consistently worse; k = 3 adds nothing over k = 2, and k = 1 is '
     f'weaker on RGB ({avg_r["geo_k1"]:+.2f}). Among single transforms, perspective is the most reliable '
     f'({avg_r["perspective"]:+.2f}, {up_r["perspective"]} sets) and crop the least ({avg_r["crop"]:+.2f}).'),
    ('Augmentation helps grayscale models more than RGB models.',
     f'The interaction difference is positive for every augmentation ({inter_vals.min():+.2f} to '
     f'{inter_vals.max():+.2f} points over 15 test sets) and stays positive without SVOX Night. On grayscale every '
     f'augmentation is consistently better on {up_g.min()}–{up_g.max()} test sets and never consistently worse '
     f'(RGB: {up_r.min()}–{up_r.max()} better, up to {dn_r.max()} worse). The largest interaction is for crop and '
     'mixed k = 1, which are weak on RGB but useful on grayscale.'),
    ('Removing colour costs a little on ordinary daytime scenes…',
     f'Without augmentation, grayscale is below RGB on Pitts30k ({none_gap["pitts30k_test"]:+.1f}), Pitts250k '
     f'({none_gap["pitts250k_test"]:+.1f}), MSLS ({none_gap["msls_val"]:+.1f}), SF-XL ({none_gap["sf_xl_small"]:+.1f}), '
     f'Tokyo 24/7 ({none_gap["tokyo247"]:+.1f}) and Nordland ({none_gap["nordland"]:+.1f}).'),
    ('…but helps greatly under strong appearance change.',
     f'SVOX Night: {gray["R1"].loc["none", "svox_night"]:.1f} % grayscale vs. {rgb["R1"].loc["none", "svox_night"]:.1f} % '
     f'RGB without augmentation, and {night_aug_gap.min():+.1f} to {night_aug_gap.max():+.1f} points with '
     f'augmentation. AmsterTime (historical, often black-and-white photos) is also better in grayscale '
     f'({none_gap["amstertime"]:+.1f} without augmentation), as is Eynsham, whose images are grayscale.'),
    ('On RGB, geometry does not fix lighting.',
     'No geometric augmentation reliably improves RGB night performance; the large night gains come from removing '
     'colour (grayscale) or, in the RandAugment reference, from colour/brightness operations. Once colour is removed, '
     f'geometric augmentation does add a further gain at night '
     f'({gray["dR1"].loc[AUGS, "svox_night"].min():+.1f} to {gray["dR1"].loc[AUGS, "svox_night"].max():+.1f} points on '
     'grayscale, consistent in all three seeds for every augmentation except rotation). This points to colour/lighting variation as the main weakness of the RGB models, motivating the '
     'photometric augmentation experiments.'),
    ('Perspective is the best augmentation for seasonal change (Nordland).',
     f'On Nordland (summer vs. winter), perspective gives the largest gain for both representations (RGB '
     f'{rgb["dR1"].loc["perspective", "nordland"]:+.1f}, grayscale {gray["dR1"].loc["perspective", "nordland"]:+.1f}), '
     'more than the mixed conditions.'),
    ('Results vary between seeds.',
     'Single runs can differ by several points on Nordland, SVOX Night, SVOX Sun and Tokyo 24/7, so conclusions rest '
     'on effects that are consistent across all three seeds rather than on single averages.'),
]
for head, body in findings:
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(head + ' '); r.bold = True
    p.add_run(body)

doc.save('results/rgb_gray_full_results.docx')
print('saved results/rgb_gray_full_results.docx, tables', tno[0])
