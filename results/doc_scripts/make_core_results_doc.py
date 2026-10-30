"""Core results document: compact tables, each followed by the trends it shows.

Run from the repo root with python-docx on PYTHONPATH:
    PYTHONPATH=~/pylib python results/doc_scripts/make_core_results_doc.py
"""
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SEEDS = [190223, 1, 2]
GEN = ['pitts30k_test', 'pitts250k_test', 'msls_val', 'sf_xl_small', 'tokyo247', 'amstertime', 'st_lucia', 'eynsham']
SEA = ['nordland', 'svox', 'svox_overcast', 'svox_rain', 'svox_snow', 'svox_sun']
NIGHT = 'svox_night'
ALL = GEN + SEA + [NIGHT]
GEO_S = ['crop', 'perspective', 'rotate', 'translate', 'shear']
GEO_M = ['geo_k1', 'geo_k2', 'geo_k3']
PHO_S = ['brightness', 'contrast', 'autocontrast', 'blur']
PHO_M = ['phot_k1', 'phot_k2', 'phot_k3']
GEO, PHO = GEO_S + GEO_M, PHO_S + PHO_M
GROUPS = [('General', GEN), ('Seas./weather', SEA), ('Night', [NIGHT])]


def load(suffix, seeds=SEEDS):
    d = pd.concat([pd.read_csv(f'results/seed{s}/test_results{suffix}.csv') for s in seeds])
    d = d[d.seed.isin(seeds)]
    p = d.pivot_table(index=['seed', 'condition'], columns='dataset', values='R1')
    g = pd.concat([p.loc[s] - p.loc[s].loc['none'] for s in seeds], keys=seeds)
    gg = g.groupby(level=1)
    return dict(abs=p.groupby(level=1).mean(), abs_sd=p.groupby(level=1).std(), g=gg.mean(),
                up=gg.agg(lambda v: bool((v > 0).all())), dn=gg.agg(lambda v: bool((v < 0).all())))


R, G = load(''), load('_gray')
f = lambda v: '0.0' if round(v, 1) == 0 else f'{v:+.1f}'
avg = lambda rep, c, ks=ALL: rep['g'].loc[c, ks].mean()
n_up = lambda rep, c: int(rep['up'].loc[c].sum())
n_dn = lambda rep, c: int(rep['dn'].loc[c].sum())

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
doc.sections[0].left_margin = doc.sections[0].right_margin = Cm(2.0)
tno, fno = [0], [0]


def border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '000000')
    b.append(el)


def caption(text, kind='Table'):
    n = tno if kind == 'Table' else fno
    n[0] += 1
    p = doc.add_paragraph()
    r = p.add_run(f'{kind} {n[0]}. '); r.bold = True
    p.add_run(text)
    if kind == 'Table':
        p.paragraph_format.keep_with_next = True


def table(header, rows, widths, font=9.5, rules=(), bold_rows=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        border(c, 'top', 12); border(c, 'bottom', 6)
    for i, cells in enumerate(rows, 1):
        rc = t.rows[i].cells
        for j, text in enumerate(cells):
            par = rc[j].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.keep_with_next = i < len(rows)
            r = par.add_run(text); r.font.size = Pt(font); r.bold = cells[0] in bold_rows
        if cells[0] in rules:
            for c in rc: border(c, 'bottom', 4)
        if i == len(rows):
            for c in rc: border(c, 'bottom', 12)


def trends(items):
    p = doc.add_paragraph(); r = p.add_run('Trends'); r.bold = True
    p.paragraph_format.space_before = Pt(6)
    for t in items:
        doc.add_paragraph(t, style='List Bullet')


# ======================================================================= title + core trends
doc.add_heading('Core results: augmentation families and input representation (MixVPR)', level=1)
doc.add_paragraph('MixVPR (ResNet-50) trained on GSV-Cities, 50 epochs, three seeds per condition, evaluated on 15 test '
                  'sets. Values are Recall@1 (%) or the change in Recall@1 (percentage points) from the '
                  'no-augmentation model with the same input, averaged over seeds. "Consistent" means the change has '
                  'the same direction in all three seeds.')
doc.add_heading('Core trends', level=2)
geo_night_max = R['g'].loc[GEO, NIGHT].max()
for t in [
    f'Removing colour is the single largest effect: SVOX Night R@1 rises from {R["abs"].loc["none", NIGHT]:.1f} % (RGB) '
    f'to {G["abs"].loc["none", NIGHT]:.1f} % (grayscale) without any augmentation.',
    'The night failure of RGB models is a colour problem, not a darkness problem: night queries are about as dark as '
    'overcast and snow queries, which RGB handles well, but carry a strong orange cast.',
    f'Photometric augmentation targets lighting change, geometric augmentation does not: on RGB at night, mixed '
    f'photometric gives {f(R["g"].loc["phot_k3", NIGHT])}, any geometric condition at most {f(geo_night_max)}.',
    'The value of an augmentation depends on the input: every geometric augmentation helps grayscale more than RGB, '
    'while colour-related photometric gains shrink once colour is removed.',
    f'Combining transforms works best: mixed conditions are the most reliable of each family (RGB average gain: '
    f'phot_k3 {f(avg(R, "phot_k3"))}, geo_k2 {f(avg(R, "geo_k2"))}).',
    f'RandAugment, which mixes geometric and colour operations, remains the strongest single policy '
    f'({f(avg(R, "randaugment"))} RGB, {f(avg(G, "randaugment"))} grayscale); grayscale + RandAugment gives the best '
    f'night result ({G["abs"].loc["randaugment", NIGHT]:.1f} %).',
]:
    doc.add_paragraph(t, style='List Bullet')

# ======================================================================= 1. night RGB vs gray
doc.add_heading('1. Night-time performance: RGB vs. grayscale', level=2)
caption('SVOX Night Recall@1 (%), mean ± standard deviation over three seeds.')
rows = []
for lab, pick in [('No augmentation', lambda rep: 'none'), ('Best geometric', lambda rep: rep['abs'].loc[GEO, NIGHT].idxmax()),
                  ('Best photometric', lambda rep: rep['abs'].loc[PHO, NIGHT].idxmax()), ('RandAugment', lambda rep: 'randaugment')]:
    cr, cg = pick(R), pick(G)
    name = lambda c: '' if c in ('none', 'randaugment') else f' ({c})'
    rows.append([lab, f'{R["abs"].loc[cr, NIGHT]:.1f} ± {R["abs_sd"].loc[cr, NIGHT]:.1f}{name(cr)}',
                 f'{G["abs"].loc[cg, NIGHT]:.1f} ± {G["abs_sd"].loc[cg, NIGHT]:.1f}{name(cg)}'])
table(['', 'RGB', 'Grayscale'], rows, [4.0, 5.2, 5.2], rules=('No augmentation',))
doc.add_picture('results/figures/night_rgb_vs_gray.png', width=Cm(9.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption('The values of Table 1 as a chart (one IEEE column; results/figures/night_rgb_vs_gray.pdf).', kind='Figure')
trends([
    f'Grayscale without augmentation ({G["abs"].loc["none", NIGHT]:.1f} %) already beats every RGB model except '
    f'RandAugment; the best RGB geometric model reaches only {R["abs"].loc[GEO, NIGHT].max():.1f} %.',
    f'On RGB, geometric augmentation barely changes night performance, photometric augmentation doubles it '
    f'({R["abs"].loc[PHO, NIGHT].max():.1f} %), and RandAugment quadruples it ({R["abs"].loc["randaugment", NIGHT]:.1f} %).',
    f'On grayscale, geometric and photometric augmentation reach similar levels ({G["abs"].loc[GEO, NIGHT].max():.1f} % vs. '
    f'{G["abs"].loc[PHO, NIGHT].max():.1f} %): once colour is removed, the family matters much less at night.',
])

# ======================================================================= 2. colour statistics
doc.add_heading('2. Why night images are hard for RGB models', level=2)
cs = pd.read_csv('results/night_analysis/colour_stats_per_image.csv').groupby('set', sort=False).mean()
r1key = {'Overcast': 'svox_overcast', 'Snow': 'svox_snow', 'Rain': 'svox_rain', 'Sun': 'svox_sun', 'Night': NIGHT}
caption('Colour statistics of SVOX images (400 per set, CIE L*a*b*) and R@1 of the RGB model without augmentation. '
        'Lightness: 0 = black, 100 = white. Cast strength: distance of the mean colour from neutral grey. '
        'b*: positive = yellow/orange, negative = blue.')
rows = []
for s in ['Database (day)', 'Overcast', 'Snow', 'Rain', 'Sun', 'Night']:
    rows.append([s, f'{cs.loc[s, "L"]:.0f}', f'{cs.loc[s, "cast"]:.1f}', f'{cs.loc[s, "b"]:+.1f}',
                 f'{R["abs"].loc["none", r1key[s]]:.1f}' if s in r1key else '—'])
table(['Set', 'Lightness', 'Cast strength', 'b*', 'RGB R@1'], rows, [3.6, 2.4, 2.8, 2.2, 2.4], rules=('Database (day)',),
      bold_rows=('Night',))
trends([
    f'Night ({cs.loc["Night", "L"]:.0f}) is barely darker than overcast ({cs.loc["Overcast", "L"]:.0f}) and snow '
    f'({cs.loc["Snow", "L"]:.0f}), yet RGB recall collapses only at night ({R["abs"].loc["none", NIGHT]:.1f} % vs. '
    f'{R["abs"].loc["none", "svox_overcast"]:.1f} % and {R["abs"].loc["none", "svox_snow"]:.1f} %): darkness alone does '
    'not explain the failure.',
    f'Night has by far the strongest colour cast ({cs.loc["Night", "cast"]:.1f}, about {cs.loc["Night", "cast"] / cs.loc["Database (day)", "cast"]:.0f}× '
    f'the daytime database) and is the only set shifted towards yellow/orange (b* {cs.loc["Night", "b"]:+.1f}); all '
    'daytime query sets are slightly blue.',
    'This matches the results in Table 1: removing colour (grayscale) or augmenting colour/intensity helps at night, '
    'changing geometry does not.',
])

# ======================================================================= 3. family x representation
doc.add_heading('3. Geometric vs. photometric, RGB vs. grayscale', level=2)
caption('Mean R@1 change from no augmentation per test-set group, for the strongest condition of each family '
        '(geo_k2, phot_k3) and the RandAugment reference. First row: absolute R@1 without augmentation.')
rows = [['No augmentation (R@1)'] + [f'{R["abs"].loc["none", ks].mean():.1f}' for _, ks in GROUPS]
        + [f'{G["abs"].loc["none", ks].mean():.1f}' for _, ks in GROUPS]]
for c, lab in [('geo_k2', 'Mixed geometric (k=2)'), ('phot_k3', 'Mixed photometric (k=3)'), ('randaugment', 'RandAugment')]:
    rows.append([lab] + [f(avg(R, c, ks)) for _, ks in GROUPS] + [f(avg(G, c, ks)) for _, ks in GROUPS])
table(['', 'RGB gen.', 'RGB seas.', 'RGB night', 'Gray gen.', 'Gray seas.', 'Gray night'], rows,
      [4.3, 1.95, 1.95, 1.95, 1.95, 1.95, 1.95], font=9, rules=('No augmentation (R@1)', 'Mixed photometric (k=3)'))
trends([
    f'On the general benchmarks both families give similar, modest gains (RGB: {f(avg(R, "geo_k2", GEN))} vs. '
    f'{f(avg(R, "phot_k3", GEN))}).',
    f'On the seasonal/weather sets photometric is ahead for RGB ({f(avg(R, "phot_k3", SEA))} vs. '
    f'{f(avg(R, "geo_k2", SEA))}), but the two are level for grayscale ({f(avg(G, "phot_k3", SEA))} vs. '
    f'{f(avg(G, "geo_k2", SEA))}).',
    f'At night the families differ most for RGB (photometric {f(R["g"].loc["phot_k3", NIGHT])} vs. geometric '
    f'{f(R["g"].loc["geo_k2", NIGHT])}) and least for grayscale ({f(G["g"].loc["phot_k3", NIGHT])} vs. '
    f'{f(G["g"].loc["geo_k2", NIGHT])}).',
    f'Grayscale costs little without augmentation on daytime sets (general {R["abs"].loc["none", GEN].mean():.1f} → '
    f'{G["abs"].loc["none", GEN].mean():.1f}) but gains greatly at night.',
])

# ======================================================================= 4. single vs mixed
doc.add_heading('4. Single vs. mixed transforms', level=2)
caption('Average R@1 change over all 15 test sets, and number of test sets improved (better) or degraded (worse) in '
        'all three seeds.')
rows = []
for c in GEO_S + GEO_M + PHO_S + PHO_M:
    rows.append([c, f(avg(R, c)), f'{n_up(R, c)} / {n_dn(R, c)}', f(avg(G, c)), f'{n_up(G, c)} / {n_dn(G, c)}'])
table(['Condition', 'RGB avg.', 'RGB better / worse', 'Gray avg.', 'Gray better / worse'], rows,
      [3.0, 2.4, 3.4, 2.4, 3.4], font=9, rules=('shear', 'geo_k3', 'blur'))
trends([
    f'Mixed conditions are the strongest and most reliable in both families: on RGB, geo_k2 ({f(avg(R, "geo_k2"))}, '
    f'{n_up(R, "geo_k2")} test sets better) beats every single geometric transform (best: perspective '
    f'{f(avg(R, "perspective"))}), and phot_k3 ({f(avg(R, "phot_k3"))}, {n_up(R, "phot_k3")} better) beats every single '
    f'photometric transform (best: contrast {f(avg(R, "contrast"))}).',
    f'Among single transforms, perspective and contrast are the most useful; crop ({f(avg(R, "crop"))}) and '
    f'autocontrast ({f(avg(R, "autocontrast"))}) barely help on RGB.',
    'On grayscale every condition improves performance on average and almost none is consistently worse anywhere; '
    'augmentation is safer to apply once colour is removed.',
])

# ======================================================================= 5. interaction
doc.add_heading('5. Does augmentation help grayscale or RGB more?', level=2)
caption('R@1 change from no augmentation for RGB and grayscale: average over all 15 test sets with the difference '
        '(Gray − RGB; positive = helps grayscale more), and on SVOX Night alone.')
rows = []
for c in GEO + PHO + ['randaugment']:
    rows.append([c, f(avg(R, c)), f(avg(G, c)), f(avg(G, c) - avg(R, c)),
                 f(R['g'].loc[c, NIGHT]), f(G['g'].loc[c, NIGHT])])
table(['Condition', 'All 15: RGB', 'All 15: Gray', 'Gray − RGB', 'Night: RGB', 'Night: Gray'], rows,
      [2.8, 2.4, 2.4, 2.4, 2.4, 2.4], font=9, rules=('geo_k3', 'phot_k3'))
trends([
    f'Geometric augmentation helps grayscale more for every condition (Gray − RGB from '
    f'{f(min(avg(G, c) - avg(R, c) for c in GEO))} to {f(max(avg(G, c) - avg(R, c) for c in GEO))}).',
    f'Photometric augmentation is mixed: contrast ({f(avg(G, "contrast") - avg(R, "contrast"))}) and blur help RGB '
    f'slightly more, brightness ({f(avg(G, "brightness") - avg(R, "brightness"))}) and phot_k2 '
    f'({f(avg(G, "phot_k2") - avg(R, "phot_k2"))}) help grayscale more.',
    f'At night (last two columns) the pattern is clearest: contrast helps RGB far more ({f(R["g"].loc["contrast", NIGHT])} vs. '
    f'{f(G["g"].loc["contrast", NIGHT])}), brightness only helps grayscale ({f(R["g"].loc["brightness", NIGHT])} vs. '
    f'{f(G["g"].loc["brightness", NIGHT])}): colour-related augmentation overlaps with removing colour, while '
    'darkness becomes the remaining problem for grayscale.',
    f'RandAugment helps RGB more ({f(avg(G, "randaugment") - avg(R, "randaugment"))}), mainly because its colour '
    'operations fix much of the RGB night problem that grayscale input already removes.',
])

# ======================================================================= 6. TrivialAugment
d = pd.concat([pd.read_csv(f'results/seed{s}/test_results.csv') for s in SEEDS])
ta_seeds = sorted(set(d[d.condition == 'trivialaugment'].seed), key=SEEDS.index)
T2 = load('', ta_seeds)
doc.add_heading('6. RandAugment vs. TrivialAugment (RGB, preliminary)', level=2)
caption(f'R@1 change from no augmentation, mean over seeds {", ".join(map(str, ta_seeds))} '
        f'({len(ta_seeds)} of 3 seeds completed).')
rows = [[lab, f(T2['g'].loc['randaugment', ks].mean()), f(T2['g'].loc['trivialaugment', ks].mean())]
        for lab, ks in [('All 15 test sets', ALL), ('General', GEN), ('Seasonal/weather', SEA), ('Nordland', ['nordland']),
                        ('SVOX Night', [NIGHT])]]
rows.append(['Training time per run', '≈ 7.0 h', '≈ 5.5 h'])
table(['', 'RandAugment', 'TrivialAugment'], rows, [4.2, 3.0, 3.0], rules=('SVOX Night',))
wins = int((T2['g'].loc['trivialaugment', ALL] > T2['g'].loc['randaugment', ALL]).sum())
trends([
    f'TrivialAugment gives smaller gains than RandAugment on {15 - wins} of 15 test sets '
    f'({f(T2["g"].loc["trivialaugment", ALL].mean())} vs. {f(T2["g"].loc["randaugment", ALL].mean())} on average).',
    f'It degrades seasonal recognition: Nordland {f(T2["g"].loc["trivialaugment", "nordland"])} (below no augmentation, '
    'in both completed seeds).',
    'It trains about 20 % faster (one operation per image instead of three), but the speed does not translate into '
    'accuracy. Preliminary until the third seed is complete.',
])

doc.save('results/core_results.docx')
print('saved results/core_results.docx, tables', tno[0])
