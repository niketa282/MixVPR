"""Answers to the research question and its sub-questions, from the results tables (results/research_questions.docx).

Run from the repo root with python-docx on PYTHONPATH:
    PYTHONPATH=<pylib> python results/doc_scripts/make_research_questions_doc.py
"""
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SEEDS = [190223, 1, 2]
GEN = ['pitts30k_test', 'pitts250k_test', 'msls_val', 'sf_xl_small', 'tokyo247', 'amstertime', 'st_lucia', 'eynsham']
WEA = ['svox_overcast', 'svox_rain', 'svox_snow', 'svox_sun']
NORD, NIGHT = 'nordland', 'svox_night'
ALL = GEN + [NORD, 'svox'] + WEA + [NIGHT]
COND = [('General (8)', GEN), ('Seasonal: Nordland', [NORD]), ('Weather (4)', WEA), ('Night: SVOX Night', [NIGHT])]
GEO_S = ['crop', 'perspective', 'rotate', 'translate', 'shear']
INT_S = ['brightness', 'contrast', 'autocontrast']
GEO_M = ['geo_k1', 'geo_k2', 'geo_k3']
PHO_M = ['phot_k1', 'phot_k2', 'phot_k3']
GEO = GEO_S + GEO_M
PHO = INT_S + ['blur'] + PHO_M
OP_GROUPS = [('Geometric (5 single ops)', GEO_S), ('Intensity (3 single ops)', INT_S), ('Degradation: blur', ['blur']),
             ('Mixed geometric (k=1-3)', GEO_M), ('Mixed photometric (k=1-3)', PHO_M)]
NAMES = {'geo_k1': 'mixed geo k=1', 'geo_k2': 'mixed geo k=2', 'geo_k3': 'mixed geo k=3', 'phot_k1': 'mixed phot k=1',
         'phot_k2': 'mixed phot k=2', 'phot_k3': 'mixed phot k=3', 'randaugment': 'RandAugment',
         'trivialaugment': 'TrivialAugment'}


def load(suffix, seeds=SEEDS):
    d = pd.concat([pd.read_csv(f'results/seed{s}/test_results{suffix}.csv') for s in SEEDS])
    d = d[d.seed.isin(seeds)]
    p = d.pivot_table(index=['seed', 'condition'], columns='dataset', values='R1')
    g = pd.concat([p.loc[s] - p.loc[s].loc['none'] for s in seeds], keys=seeds)
    return dict(abs=p.groupby(level=1).mean(), g=g.groupby(level=1).mean())


R, G = load(''), load('_gray')
R12 = load('', [1, 2])  # TrivialAugment has seeds 1 and 2
f = lambda v: '0.0' if round(v, 1) == 0 else f'{v:+.1f}'
a = lambda v: f'{v:.1f}'
gain = lambda rep, c, ks: rep['g'].loc[c, ks].mean()
grp = lambda rep, cs, ks: rep['g'].loc[cs, ks].mean(axis=1).mean()
n_up = lambda rep, c: int((rep['g'].loc[c, ALL] > 0).sum())


def best(rep, conds, ks):
    s = rep['g'].loc[conds, ks].mean(axis=1)
    return s.idxmax(), s.max()


doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
doc.sections[0].left_margin = doc.sections[0].right_margin = Cm(2.0)
tno = [0]


def border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '000000')
    b.append(el)


def caption(text):
    tno[0] += 1
    p = doc.add_paragraph()
    r = p.add_run(f'Table {tno[0]}. '); r.bold = True
    p.add_run(text)
    p.paragraph_format.keep_with_next = True


def table(header, rows, widths, font=9.5, rules=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        border(c, 'top', 12); border(c, 'bottom', 6)
    for i, cells in enumerate(rows, 1):
        rc = t.rows[i].cells
        for j, text in enumerate(cells):
            par = rc[j].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.keep_with_next = i < len(rows)
            r = par.add_run(text); r.font.size = Pt(font)
        if i in rules:
            for c in rc: border(c, 'bottom', 4)
        if i == len(rows):
            for c in rc: border(c, 'bottom', 12)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def answer(text):
    p = doc.add_paragraph()
    r = p.add_run('Answer. '); r.bold = True
    p.add_run(text)


def bullets(items):
    for t in items:
        doc.add_paragraph(t, style='List Bullet')


# ================================================================================================ title + setup
doc.add_heading('Research questions: answers from the results', level=1)
p = doc.add_paragraph()
r = p.add_run('Main question. '); r.bold = True
p.add_run('How do augmentation operations and policies affect the robustness of visual place recognition, and does '
          'this effect depend on the input representation (RGB vs. grayscale)?')
doc.add_paragraph('Setup: MixVPR (ResNet-50) trained on GSV-Cities for 50 epochs with each augmentation condition, '
                  'once with RGB and once with grayscale input, three seeds per condition. Evaluated on 15 test sets: '
                  '8 general benchmarks (Pitts30k, Pitts250k, MSLS-val, SF-XL small, Tokyo 24/7, AmsterTime, '
                  'St Lucia, Eynsham), seasonal change (Nordland, summer database vs. winter queries), weather '
                  '(SVOX overcast, rain, snow, sun), night (SVOX Night) and SVOX (all). Values are Recall@1 (%) or the '
                  'gain in Recall@1 (percentage points) over the no-augmentation model with the same input, '
                  'averaged over seeds.')

# ================================================================================================ SQ1: policies
doc.add_heading('SQ1. Which augmentation policies give the best recall on RGB versus grayscale input?', level=2)
caption('Augmentation policies: average R@1 gain over no augmentation, and absolute R@1 on SVOX Night. '
        'TrivialAugment was trained on RGB with seeds 1 and 2; its RandAugment comparison row uses the same seeds.')
rows = []
for lab, rep, c in [('RandAugment, RGB', R, 'randaugment'), ('RandAugment, grayscale', G, 'randaugment'),
                    ('RandAugment, RGB (seeds 1, 2)', R12, 'randaugment'),
                    ('TrivialAugment, RGB (seeds 1, 2)', R12, 'trivialaugment')]:
    rows.append([lab, f(gain(rep, c, ALL)), f(gain(rep, c, GEN)), f(gain(rep, c, [NORD])), f(gain(rep, c, WEA)),
                 f(gain(rep, c, [NIGHT])), a(rep['abs'].loc[c, NIGHT]), f'{n_up(rep, c)}/15'])
table(['Policy, input', 'All 15', 'General', 'Nordland', 'Weather', 'Night', 'Night R@1', 'Sets improved'], rows,
      [5.0, 1.4, 1.5, 1.6, 1.5, 1.4, 1.7, 1.9], font=9, rules=(2,))

ra_r, ra_g = R['abs'].loc['randaugment'], G['abs'].loc['randaugment']
day = [k for k in ALL if k != NIGHT]
answer(f'RandAugment gives the best recall of the policies for both inputs. It raises average R@1 over all 15 test '
       f'sets by {f(gain(R, "randaugment", ALL))} points on RGB and {f(gain(G, "randaugment", ALL))} on grayscale. '
       f'With RandAugment, grayscale reaches the highest overall recall ({a(ra_g[ALL].mean())} % vs. '
       f'{a(ra_r[ALL].mean())} % for RGB), because it is far better at night ({a(ra_g[NIGHT])} % vs. '
       f'{a(ra_r[NIGHT])} %); on the 14 daytime sets RGB is slightly higher ({a(ra_r[day].mean())} % vs. '
       f'{a(ra_g[day].mean())} %). TrivialAugment on RGB gives smaller gains than RandAugment '
       f'({f(gain(R12, "trivialaugment", ALL))} vs. {f(gain(R12, "randaugment", ALL))} on the same seeds) but '
       f'trains about 20 % faster (about 5.5 h vs. 7.0 h).')
bullets([
    f'RandAugment\'s gain is largest at night: {f(gain(R, "randaugment", [NIGHT]))} on RGB '
    f'({a(R["abs"].loc["none", NIGHT])} % to {a(ra_r[NIGHT])} %) and {f(gain(G, "randaugment", [NIGHT]))} on '
    f'grayscale ({a(G["abs"].loc["none", NIGHT])} % to {a(ra_g[NIGHT])} %). The grayscale gain is smaller because '
    f'its starting point is much higher.',
    f'On the general benchmarks the policy gains are small and similar for both inputs '
    f'({f(gain(R, "randaugment", GEN))} RGB, {f(gain(G, "randaugment", GEN))} grayscale).',
    f'TrivialAugment lowers Nordland R@1 by {abs(gain(R12, "trivialaugment", [NORD])):.1f} points, whereas '
    f'RandAugment raises it by {gain(R12, "randaugment", [NORD]):.1f} on the same seeds. TrivialAugment is closest to '
    f'RandAugment on the general benchmarks ({f(gain(R12, "trivialaugment", GEN))} vs. '
    f'{f(gain(R12, "randaugment", GEN))}) and gives the largest gain of all conditions on AmsterTime '
    f'({f(R12["g"].loc["trivialaugment", "amstertime"])} vs. {f(R12["g"].loc["randaugment", "amstertime"])}).',
    f'RandAugment improves {n_up(G, "randaugment")}/15 test sets with grayscale and {n_up(R, "randaugment")}/15 with '
    f'RGB. On individual test sets a targeted geometric or photometric condition beats it on 6/15 grayscale sets '
    f'(e.g. Nordland: mixed photometric {f(G["g"].loc["phot_k3", NORD])} vs. {f(G["g"].loc["randaugment", NORD])}) '
    f'and 3/15 RGB sets.',
])

# ================================================================================================ SQ2: op groups
doc.add_heading('SQ2. Which operation groups contribute most to robustness under each representation?', level=2)
caption('Average R@1 gain per operation group (mean over the operations in the group). Geometric: crop, perspective, '
        'rotate, translate, shear. Intensity: brightness, contrast, autocontrast. Degradation: Gaussian blur. Mixed: '
        'k operations drawn at random per image from the group (photometric = intensity + blur).')
rows = []
for lab, cs in OP_GROUPS:
    for rl, rep in (('RGB', R), ('Gray', G)):
        rows.append([f'{lab}, {rl}', f(grp(rep, cs, ALL)), f(grp(rep, cs, GEN)), f(grp(rep, cs, [NORD])),
                     f(grp(rep, cs, WEA)), f(grp(rep, cs, [NIGHT]))])
table(['Group, input', 'All 15', 'General', 'Nordland', 'Weather', 'Night'], rows, [6.2, 1.6, 1.6, 1.7, 1.6, 1.5],
      font=9, rules=(2, 4, 6, 8))

caption('Single operations: average R@1 gain over all 15 sets, gain on SVOX Night and number of test sets improved.')
rows = []
for c in GEO_S + INT_S + ['blur']:
    rows.append([c, f(gain(R, c, ALL)), f(gain(R, c, [NIGHT])), f'{n_up(R, c)}/15',
                 f(gain(G, c, ALL)), f(gain(G, c, [NIGHT])), f'{n_up(G, c)}/15'])
table(['Operation', 'RGB all 15', 'RGB night', 'RGB improved', 'Gray all 15', 'Gray night', 'Gray improved'], rows,
      [2.8, 2.0, 2.0, 2.1, 2.0, 2.0, 2.1], font=9, rules=(5, 8))

geo_r, geo_g = grp(R, GEO_S, ALL), grp(G, GEO_S, ALL)
answer(f'Under grayscale, geometric operations contribute most: the single geometric operations raise average R@1 by '
       f'{f(geo_g)} points, about three times as much as on RGB ({f(geo_r)}), and more than intensity '
       f'({f(grp(G, INT_S, ALL))}) or blur ({f(grp(G, ["blur"], ALL))}). Under RGB, intensity operations contribute '
       f'most, and almost entirely through night: intensity operations give {f(grp(R, INT_S, [NIGHT]))} at night '
       f'(contrast alone {f(R["g"].loc["contrast", NIGHT])}), geometric operations {f(grp(R, GEO_S, [NIGHT]))}. In '
       f'daytime the groups contribute similarly on RGB (general: geometric {f(grp(R, GEO_S, GEN))}, intensity '
       f'{f(grp(R, INT_S, GEN))}, blur {f(grp(R, ["blur"], GEN))}).')
bullets([
    'Every geometric operation helps grayscale more than RGB (Table 3). With grayscale each single geometric '
    f'operation improves at least {min(n_up(G, c) for c in GEO_S)}/15 test sets; with RGB the range is '
    f'{min(n_up(R, c) for c in GEO_S)}-{max(n_up(R, c) for c in GEO_S)}/15.',
    f'Geometric operations also help grayscale at night ({f(grp(G, GEO_S, [NIGHT]))} on average, shear '
    f'{f(G["g"].loc["shear", NIGHT])}), while on RGB they leave night performance essentially unchanged.',
    f'The intensity group interacts with colour: contrast gives {f(R["g"].loc["contrast", NIGHT])} at night on RGB '
    f'but {f(G["g"].loc["contrast", NIGHT])} on grayscale; brightness gives {f(R["g"].loc["brightness", NIGHT])} on '
    f'RGB and {f(G["g"].loc["brightness", NIGHT])} on grayscale.',
    f'Blur (degradation) contributes most on Nordland ({f(R["g"].loc["blur", NORD])} RGB, '
    f'{f(G["g"].loc["blur", NORD])} grayscale) and little on the general benchmarks.',
    f'Combining operations is stronger than single operations in every group and representation. Mixed photometric is '
    f'the strongest group on RGB ({f(grp(R, PHO_M, ALL))}; k=3: {f(gain(R, "phot_k3", ALL))}, night '
    f'{f(R["g"].loc["phot_k3", NIGHT])}); on grayscale mixed geometric ({f(grp(G, GEO_M, ALL))}) and mixed '
    f'photometric ({f(grp(G, PHO_M, ALL))}) are similar and improve all 15 test sets.',
])

# ================================================================================================ SQ3: conditions
doc.add_heading('SQ3. Under which conditions do these effects appear?', level=2)
caption('Baseline R@1 (no augmentation) per condition, RGB vs. grayscale.')
rows = []
for lab, ks in COND:
    vr, vg = R['abs'].loc['none', ks].mean(), G['abs'].loc['none', ks].mean()
    rows.append([lab, a(vr), a(vg), f(vg - vr)])
table(['Condition', 'RGB', 'Grayscale', 'Gray - RGB'], rows, [5.0, 2.2, 2.4, 2.4], font=9.5)

caption('Largest R@1 gain per condition for each family (condition in brackets) and for RandAugment.')
rows = []
for lab, ks in COND:
    for rl, rep in (('RGB', R), ('Gray', G)):
        cg, vg = best(rep, GEO, ks)
        cp, vp = best(rep, PHO, ks)
        rows.append([f'{lab}, {rl}', f'{f(vg)} ({NAMES.get(cg, cg)})', f'{f(vp)} ({NAMES.get(cp, cp)})',
                     f(gain(rep, 'randaugment', ks))])
table(['Condition, input', 'Best geometric', 'Best photometric', 'RandAugment'], rows, [5.4, 4.1, 4.4, 2.3],
      font=9, rules=(2, 4, 6))

b = lambda rep, fam, ks: best(rep, fam, ks)[1]
answer('The effects appear most strongly at night, then under seasonal change and weather, and are smallest on the '
       'general urban benchmarks. The dependence on the input representation is also concentrated at night.')
bullets([
    f'Night: the input representation is the largest single effect. Without augmentation grayscale reaches '
    f'{a(G["abs"].loc["none", NIGHT])} % R@1 vs. {a(R["abs"].loc["none", NIGHT])} % for RGB; night images carry a '
    f'strong orange colour cast that grayscale removes. On RGB only photometric augmentation helps at night '
    f'({f(b(R, PHO, [NIGHT]))} vs. at most {f(b(R, GEO, [NIGHT]))} for geometric); on grayscale both families help '
    f'({f(b(G, GEO, [NIGHT]))} geometric, {f(b(G, PHO, [NIGHT]))} photometric).',
    f'Seasonal change (Nordland): the largest daytime gains. Photometric augmentation is strongest for both inputs '
    f'({f(b(R, PHO, [NORD]))} RGB, {f(b(G, PHO, [NORD]))} grayscale) and exceeds RandAugment on grayscale; '
    f'perspective is the strongest geometric operation ({f(R["g"].loc["perspective", NORD])} RGB, '
    f'{f(G["g"].loc["perspective", NORD])} grayscale). Grayscale starts '
    f'{abs(G["abs"].loc["none", NORD] - R["abs"].loc["none", NORD]):.1f} points lower here.',
    f'Weather (SVOX overcast, rain, snow, sun): moderate gains. On RGB photometric is stronger '
    f'({f(b(R, PHO, WEA))} vs. {f(b(R, GEO, WEA))} geometric); on grayscale geometric is stronger '
    f'({f(b(G, GEO, WEA))} vs. {f(b(G, PHO, WEA))}). Sun benefits most and overcast least.',
    f'General benchmarks: small gains (best family condition {f(b(R, GEO + PHO, GEN))} RGB, '
    f'{f(b(G, GEO + PHO, GEN))} grayscale) and the two inputs start within '
    f'{abs(G["abs"].loc["none", GEN].mean() - R["abs"].loc["none", GEN].mean()):.1f} points of each other. '
    f'Grayscale is ahead on AmsterTime ({f(G["abs"].loc["none", "amstertime"] - R["abs"].loc["none", "amstertime"])})'
    f' and behind on Tokyo 24/7 '
    f'({f(G["abs"].loc["none", "tokyo247"] - R["abs"].loc["none", "tokyo247"])}).',
])

# ================================================================================================ main answer
doc.add_heading('Main question: summary answer', level=2)
doc.add_paragraph(
    f'Augmentation improves VPR robustness for both inputs, and the size and source of the improvement depend on the '
    f'input representation. Geometric operations are the main contributor for grayscale input (average gain '
    f'{f(geo_g)} vs. {f(geo_r)} on RGB), while intensity/photometric operations are the main contributor for RGB, '
    f'mostly at night. Combining operations is stronger than single operations, and the RandAugment policy gives the '
    f'largest average gain for both inputs ({f(gain(R, "randaugment", ALL))} RGB, '
    f'{f(gain(G, "randaugment", ALL))} grayscale), ahead of TrivialAugment on RGB. The effects are concentrated at '
    f'night, where the choice of input representation matters more than any augmentation on its own: grayscale '
    f'without augmentation ({a(G["abs"].loc["none", NIGHT])} %) outperforms every RGB model except RandAugment, and '
    f'grayscale with RandAugment is the best model overall ({a(ra_g[NIGHT])} % at night).')

out = 'results/research_questions.docx'
doc.save(out)
print('saved', out)
