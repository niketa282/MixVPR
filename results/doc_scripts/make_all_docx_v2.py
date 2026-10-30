"""Two 3-seed results documents over all 15 test sets: geometric only, and geometric + RandAugment.
Run from the MixVPR folder with python-docx on PYTHONPATH."""
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SEEDS = [190223, 1, 2]
SINGLE = ['crop', 'perspective', 'rotate', 'translate', 'shear']
MIXED = ['geo_k1', 'geo_k2', 'geo_k3']
GEO = ['none'] + SINGLE + MIXED
AUGS = SINGLE + MIXED
TRAIN_HOURS = {'geo_k1': 6.4, 'geo_k2': 7.7, 'geo_k3': 9.3}  # mean wall-clock per run (sacct)
METRICS = ['R1', 'R5', 'R10']
GREY = RGBColor(0x5f, 0x5e, 0x58)

# (key, short name, long name, what it tests)
GROUP_A = [
    ('pitts30k_test', 'Pitts30k', 'Pitts30k-test', 'Pittsburgh street view; same city as the validation set'),
    ('pitts250k_test', 'Pitts250k', 'Pitts250k-test', 'larger Pittsburgh benchmark'),
    ('msls_val', 'MSLS', 'MSLS-val (all queries)', 'dashcam images from several cities, mixed conditions'),
    ('sf_xl_small', 'SF-XL', 'SF-XL small (test v1)', 'San Francisco; crowd-sourced query photos'),
    ('tokyo247', 'Tokyo 24/7', 'Tokyo 24/7', 'Tokyo; day, sunset and night queries'),
    ('amstertime', 'AmsterTime', 'AmsterTime', 'historical vs. modern photos of Amsterdam'),
    ('st_lucia', 'St Lucia', 'St Lucia', 'suburban driving at different times of day'),
    ('eynsham', 'Eynsham', 'Eynsham', 'rural/suburban roads; grayscale camera'),
]
GROUP_B = [
    ('nordland', 'Nordland', 'Nordland (summer database, winter queries)', 'rural railway; summer vs. winter'),
    ('svox', 'SVOX', 'SVOX (all queries)', 'Oxford; all query conditions'),
    ('svox_night', 'Night', 'SVOX Night', 'night queries'),
    ('svox_overcast', 'Overcast', 'SVOX Overcast', 'overcast queries'),
    ('svox_rain', 'Rain', 'SVOX Rain', 'rain queries'),
    ('svox_snow', 'Snow', 'SVOX Snow', 'snow queries'),
    ('svox_sun', 'Sun', 'SVOX Sun', 'strong sunlight queries'),
]
SETS = GROUP_A + GROUP_B
SHORT = {k: s for k, s, _, _ in SETS}
LONG = {k: l for k, _, l, _ in SETS}

d = pd.concat([pd.read_csv(f'results/seed{s}/test_results.csv') for s in SEEDS])
assert not d.duplicated(['seed', 'condition', 'dataset']).any()
counts = d.groupby('dataset')[['num_db', 'num_queries']].first()
agg = d.groupby(['condition', 'dataset'])[METRICS].agg(['mean', 'std'])

base = d[d.condition == 'none'].set_index(['seed', 'dataset'])['R1']
dd = d.set_index(['seed', 'dataset'])
gain = (dd['R1'] - base.reindex(dd.index)).rename('g').to_frame().assign(condition=dd.condition).reset_index()
gg = gain.groupby(['condition', 'dataset'])['g']
gmean, gstd = gg.mean(), gg.std()
gpos, gneg = gg.agg(lambda v: int((v > 0).sum())), gg.agg(lambda v: int((v < 0).sum()))


def G(c, k):
    return gmean[(c, k)], gstd[(c, k)], gpos[(c, k)], gneg[(c, k)]


def consistent_up(c):
    return [k for k, *_ in SETS if gpos[(c, k)] == 3]


def consistent_down(c):
    return [k for k, *_ in SETS if gneg[(c, k)] == 3]


def names(keys):
    return ', '.join(SHORT[k] if k not in ('svox', 'svox_night', 'svox_overcast', 'svox_rain', 'svox_snow', 'svox_sun')
                     else ('SVOX' if k == 'svox' else 'SVOX ' + SHORT[k]) for k in keys)


def border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '000000')
    b.append(el)


def build(with_ra, out):
    conds = GEO + (['randaugment'] if with_ra else [])
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    doc.sections[0].left_margin = doc.sections[0].right_margin = Cm(2.2)
    tno = [0]

    def section(landscape):
        s = doc.add_section(WD_SECTION.NEW_PAGE)
        w, h = s.page_width, s.page_height
        if (landscape and w < h) or (not landscape and w > h):
            s.page_width, s.page_height = h, w
        s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        s.left_margin = s.right_margin = Cm(1.8 if landscape else 2.2)

    def caption(text):
        tno[0] += 1
        p = doc.add_paragraph()
        r = p.add_run(f'Table {tno[0]}. '); r.bold = True
        p.add_run(text)
        p.paragraph_format.keep_with_next = True

    def table(header, rows, widths, font=9.0, rules=()):
        tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        for row in tbl.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
        for j, h in enumerate(header):
            p = tbl.rows[0].cells[j].paragraphs[0]
            r = p.add_run(h); r.bold = True; r.font.size = Pt(font)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            border(tbl.rows[0].cells[j], 'top', 12); border(tbl.rows[0].cells[j], 'bottom', 6)
        for ri, (label, cells) in enumerate(rows, 1):
            rc = tbl.rows[ri].cells
            r = rc[0].paragraphs[0].add_run(label); r.font.size = Pt(font)
            for j, parts in enumerate(cells, 1):
                p = rc[j].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if isinstance(parts, str) else WD_ALIGN_PARAGRAPH.CENTER
                if isinstance(parts, str):
                    parts = [(parts, False, False)]
                for text, bold, small in parts:
                    r = p.add_run(text); r.bold = bold
                    r.font.size = Pt(font - 1.5 if small else font)
                    if small:
                        r.font.color.rgb = GREY
            if label in rules:
                for c in rc: border(c, 'bottom', 4)
            if ri == len(rows):
                for c in rc: border(c, 'bottom', 12)
            for c in rc:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.keep_with_next = ri < len(rows)
        doc.add_paragraph()

    def best_geo(k, m):
        return max(agg.loc[(c, k), (m, 'mean')] for c in GEO)

    def ms(c, k, m, bold_best=True):
        mean, std = agg.loc[(c, k), (m, 'mean')], agg.loc[(c, k), (m, 'std')]
        b = bold_best and c in GEO and round(mean, 2) == round(best_geo(k, m), 2)
        return [(f'{mean:.2f}', b, False), (f' ± {std:.2f}', False, True)]

    rules = ('none', 'shear', 'geo_k3') if with_ra else ('none', 'shear')

    # ---------- front matter ----------
    title = 'Single and mixed geometric augmentation for MixVPR: results on 15 test sets (three seeds)'
    if with_ra:
        title = ('Single and mixed geometric augmentation and RandAugment for MixVPR: results on 15 test sets '
                 '(three seeds)')
    doc.add_heading(title, level=1)
    doc.add_heading('Setup', level=2)
    conds_txt = ('none (no augmentation); five single geometric augmentations (crop, perspective, rotate, '
                 'translate, shear); and three mixed geometric conditions, geo_k1, geo_k2 and geo_k3, which apply '
                 'exactly 1, 2 or 3 of the five single transforms to each training image, chosen at random and in '
                 'random order (same transform settings as the single conditions).')
    if with_ra:
        conds_txt += (' RandAugment, the original MixVPR augmentation policy, is included as a reference. It '
                      'mixes geometric and colour operations, so it is not a purely geometric condition.')
    for t in [
        'Model: MixVPR (ResNet-50 backbone, 4096-dimensional descriptor), RGB input, 320 × 320 images.',
        'Training: GSV-Cities, 50 epochs, identical settings for every condition. Each condition was trained '
        'three times with different random seeds (190223, 1, 2).',
        'Checkpoint: for each run, the epoch with the highest Pitts30k-val Recall@1.',
        'Conditions: ' + conds_txt,
        'Evaluation: a retrieved database image is correct if it lies within 25 m of the query (±10 frames '
        'for Nordland; the single matching image for AmsterTime). Queries with no database image within '
        '25 m count as failures.',
    ]:
        doc.add_paragraph(t, style='List Bullet')

    caption('Test sets.')
    rows = [(LONG[k], [what, f'{counts.loc[k, "num_db"]:,}', f'{counts.loc[k, "num_queries"]:,}'])
            for k, _, _, what in SETS]
    table(['Test set', 'What it tests', 'Database', 'Queries'], rows, [4.6, 6.6, 2.2, 2.2], font=9,
          rules=('Eynsham',))
    for t in [
        'Tokyo 24/7 has only 315 queries, so one query changes R@1 by 0.32 points; the SVOX subsets have '
        '820–940 queries (about 0.12 points per query).',
        'MSLS uses all 11,084 validation queries, not the 740-query subset used in published MSLS-val '
        'results, so absolute MSLS numbers are not comparable with papers.',
        'Eynsham images are grayscale; they are fed to the RGB models as three identical channels.',
    ]:
        doc.add_paragraph(t, style='List Bullet')

    doc.add_heading('How to read the tables', level=2)
    for t in [
        'Recall@K (R@K): the percentage of queries for which at least one correct place is among the top K '
        'retrieved images. R@1 is the strictest and most commonly reported.',
        'Values are mean ± standard deviation over the three seeds (sample standard deviation, n = 3). The '
        'standard deviation shows how much a result changes just by re-running training with another seed.',
        'Gain tables give the change in R@1 relative to "none", computed within each seed and then averaged. '
        'The bracket shows in how many of the three seeds the change was positive, e.g. (3/3).',
        'In the gain tables, bold marks a change with the same sign in all three seeds and larger than its '
        'standard deviation: the most reliable effects.',
        'In the recall tables, bold marks the best mean among none and the geometric augmentations' +
        (' (RandAugment, below the second rule, is not included in the bolding).' if with_ra else '.'),
    ]:
        doc.add_paragraph(t, style='List Bullet')

    # ---------- overview ----------
    doc.add_heading('Overview', level=2)
    caption('Overview over all 15 test sets: number of test sets on which the condition changed R@1 relative to '
            'no augmentation in the same direction in all three seeds, and the mean R@1 change averaged over the '
            '15 test sets (each test set weighted equally).')
    rows = []
    for c in conds[1:]:
        up_, dn_ = len(consistent_up(c)), len(consistent_down(c))
        avg = sum(G(c, k)[0] for k, *_ in SETS) / len(SETS)
        rows.append((c, [str(up_), str(dn_), f'{avg:+.2f}']))
    table(['Condition', 'Better in 3/3 seeds', 'Worse in 3/3 seeds', 'Mean R@1 change'], rows,
          [3.0, 3.6, 3.6, 3.6], font=9.5, rules=('shear', 'geo_k3') if with_ra else ('shear',))

    # ---------- observations (computed) ----------
    doc.add_heading('Key observations', level=2)
    obs = []
    up = {c: consistent_up(c) for c in AUGS}
    down = {c: consistent_down(c) for c in AUGS}
    avg = {c: sum(G(c, k)[0] for k, *_ in SETS) / len(SETS) for c in AUGS}
    k1, k2, k3 = ({k: G(c, k) for k, *_ in SETS} for c in MIXED)
    p_ = {k: G('perspective', k) for k, *_ in SETS}
    obs.append(f'Mixing geometric transforms works best. geo_k2 and geo_k3 improve R@1 in all three seeds on '
               f'{len(up["geo_k2"])} and {len(up["geo_k3"])} of 15 test sets and are never consistently worse than '
               f'no augmentation; averaged over all test sets they gain {avg["geo_k2"]:+.2f} and {avg["geo_k3"]:+.2f} '
               f'points, ahead of the best single transform, perspective ({avg["perspective"]:+.2f}, '
               f'{len(up["perspective"])} test sets).')
    obs.append(f'geo_k2 gives the largest geometric gains on SVOX Sun ({k2["svox_sun"][0]:+.1f}), SF-XL '
               f'({k2["sf_xl_small"][0]:+.1f}) and MSLS ({k2["msls_val"][0]:+.1f}), and is the only geometric '
               f'condition that improves Tokyo 24/7 in all three seeds ({k2["tokyo247"][0]:+.1f}).')
    obs.append(f'Two transforms per image are enough: geo_k3 is about level with geo_k2 (no test set where it is '
               f'clearly better) but takes longer to train (about {TRAIN_HOURS["geo_k3"]:.1f} h per run vs. '
               f'{TRAIN_HOURS["geo_k2"]:.1f} h). geo_k1, one random transform per image, is weaker: consistent gains '
               f'on only {len(up["geo_k1"])} test sets, average {avg["geo_k1"]:+.2f}, and worse than no augmentation '
               f'on SVOX Night in all three seeds ({k1["svox_night"][0]:+.1f}). Variety alone helps less than '
               f'applying more than one transform.')
    obs.append(f'On Nordland the best single transform beats the mixed conditions: perspective gains '
               f'{p_["nordland"][0]:+.1f} ± {p_["nordland"][1]:.1f} (3/3 seeds), the mixed conditions '
               f'{min(k1["nordland"][0], k2["nordland"][0], k3["nordland"][0]):+.1f} to '
               f'{max(k1["nordland"][0], k2["nordland"][0], k3["nordland"][0]):+.1f}, not consistently.')
    order = sorted(SINGLE, key=lambda c: -len(up[c]))
    obs.append('Single transforms, test sets (out of 15) improved in all three seeds: ' +
               '; '.join(f'{c} {len(up[c])}' for c in order) + '; worse in all three seeds: ' +
               '; '.join(f'{c} {len(down[c])}' for c in order) + '. Perspective is the most reliable single '
               f'transform (never consistently worse); crop is the weakest and is worse on Tokyo 24/7 '
               f'({G("crop", "tokyo247")[0]:+.1f}) and Pitts30k; translate is also worse on Tokyo 24/7 '
               f'({G("translate", "tokyo247")[0]:+.1f}).')
    night = [G(c, 'svox_night')[0] for c in AUGS]
    nsd = [G(c, 'svox_night')[1] for c in AUGS]
    tokyo_up = [c for c in AUGS if gpos[(c, 'tokyo247')] == 3]
    obs.append(f'No geometric condition reliably helps with night images: on SVOX Night the mean change ranges '
               f'from {min(night):+.1f} to {max(night):+.1f} points with standard deviations of '
               f'{min(nsd):.0f}–{max(nsd):.0f} points; on Tokyo 24/7 (day/sunset/night) only '
               f'{", ".join(tokyo_up) if tokyo_up else "none"} improves in all three seeds.')
    obs.append('St Lucia (≈99.7 % R@1), SVOX (all) and SVOX Overcast are at or near their ceiling; no condition '
               'changes them by more than about one point. Eynsham changes by less than half a point.')
    obs.append(f'Seed-to-seed variation differs widely between test sets: the baseline standard deviation is '
               f'{agg.loc[("none", "nordland"), ("R1", "std")]:.1f} points on Nordland, '
               f'{agg.loc[("none", "svox_sun"), ("R1", "std")]:.1f} on SVOX Sun and '
               f'{agg.loc[("none", "svox_night"), ("R1", "std")]:.1f} on SVOX Night, but under 0.4 on Pitts30k, '
               'Pitts250k, MSLS and SVOX (all). On the low-variance sets, gains of a few tenths of a point can '
               'still be consistent; on the high-variance sets, only large gains are trustworthy.')
    if with_ra:
        ra_up = [k for k, *_ in SETS if gpos[('randaugment', k)] == 3]
        ra = {k: G('randaugment', k) for k, *_ in SETS}
        ra_avg = sum(ra[k][0] for k, *_ in SETS) / len(SETS)
        obs.append(f'RandAugment (reference) improves R@1 in all three seeds on {len(ra_up)} of 15 test sets '
                   f'(average {ra_avg:+.2f}, against {avg["geo_k2"]:+.2f} for geo_k2). Its advantage comes mainly '
                   f'from lighting change: SVOX Night {ra["svox_night"][0]:+.1f} ± {ra["svox_night"][1]:.1f}, '
                   f'SVOX Sun {ra["svox_sun"][0]:+.1f}, Tokyo 24/7 {ra["tokyo247"][0]:+.1f}, SVOX Rain '
                   f'{ra["svox_rain"][0]:+.1f}, which points to its colour/brightness operations. On several '
                   f'general benchmarks the mixed geometric conditions come close or match it (SF-XL: geo_k2 '
                   f'{k2["sf_xl_small"][0]:+.1f} vs. {ra["sf_xl_small"][0]:+.1f}); on Nordland perspective matches '
                   f'it ({p_["nordland"][0]:+.1f} vs. {ra["nordland"][0]:+.1f}); on the grayscale Eynsham set '
                   f'RandAugment gives no gain ({ra["eynsham"][0]:+.2f}).')
    for t in obs:
        doc.add_paragraph(t, style='List Bullet')

    # ---------- summary tables (landscape) ----------
    for gi, (group, gname) in enumerate([(GROUP_A, 'General benchmarks'),
                                         (GROUP_B, 'Seasonal and weather/lighting change')]):
        section(True)
        doc.add_heading(f'Summary: {gname}', level=2)
        hdr = ['Condition'] + [s for _, s, _, _ in group]
        w = [2.4] + [(25.9 - 2.4) / len(group)] * len(group)
        caption(f'Recall@1 (%) on the {gname.lower()}, mean ± standard deviation over three seeds. Bold: '
                'best mean among none and the geometric augmentations.')
        table(hdr, [(c, [ms(c, k, 'R1') for k, *_ in group]) for c in conds], w, font=8.5, rules=rules)
        caption(f'Change in Recall@1 relative to no augmentation (percentage points) on the {gname.lower()}: '
                'mean ± standard deviation of the per-seed difference, and number of seeds with a positive '
                'change. Bold: same sign in all three seeds and larger than the standard deviation.')
        rows = []
        for c in conds[1:]:
            cells = []
            for k, *_ in group:
                m, s, pos, neg = G(c, k)
                strong = (pos == 3 or neg == 3) and abs(m) > s
                cells.append([(f'{m:+.2f}', strong, False), (f' ± {s:.2f}', False, True),
                              (f' ({pos}/3)', False, True)])
            rows.append((c, cells))
        table(hdr, rows, w, font=8.5, rules=('shear', 'geo_k3') if with_ra else ('shear',))

    # ---------- per test set ----------
    section(False)
    doc.add_heading('Results per test set: Recall@1, @5 and @10', level=2)
    for k, s, l, what in SETS:
        doc.add_heading(l, level=3)
        caption(f'Recall@K (%) on {l} ({what}): {counts.loc[k, "num_db"]:,} database images, '
                f'{counts.loc[k, "num_queries"]:,} queries. Mean ± standard deviation over three seeds; last '
                'column: R@1 change vs. none (seeds positive).')
        rows = []
        for c in conds:
            cells = [ms(c, k, m) for m in METRICS]
            if c == 'none':
                cells.append('—')
            else:
                m, sd, pos, neg = G(c, k)
                cells.append([(f'{m:+.2f}', False, False), (f' ({pos}/3)', False, True)])
            rows.append((c, cells))
        table(['Condition', 'R@1', 'R@5', 'R@10', 'R@1 change'], rows, [2.8, 3.1, 3.1, 3.1, 2.8], font=9.5,
              rules=rules)

    # ---------- appendix ----------
    doc.add_heading('Appendix: Recall@1 for each seed', level=2)
    doc.add_paragraph('The individual runs behind the means above, so the seed-to-seed variation is visible.')
    r1 = d.pivot_table(index=['dataset', 'condition'], columns='seed', values='R1')
    for k, s, l, _ in SETS:
        caption(f'Recall@1 (%) per seed on {l}.')
        rows = [(c, [[(f'{r1.loc[(k, c), sd]:.2f}', False, False)] for sd in SEEDS] + [ms(c, k, 'R1', False)])
                for c in conds]
        table(['Condition'] + [f'Seed {sd}' for sd in SEEDS] + ['Mean ± std'], rows,
              [2.8, 2.6, 2.6, 2.6, 3.4], font=9.5, rules=rules)

    doc.save(out)
    print('saved', out, 'tables:', tno[0])


build(False, 'results/geometric_results_all_datasets.docx')
build(True, 'results/geometric_results_all_datasets_with_randaugment.docx')
