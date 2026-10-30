"""Grayscale results + RGB comparison document. Run from MixVPR with python-docx on PYTHONPATH."""
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

S = [190223, 1, 2]
SINGLE = ['crop', 'perspective', 'rotate', 'translate', 'shear']
MIXED = ['geo_k1', 'geo_k2', 'geo_k3']
CONDS = ['none'] + SINGLE + MIXED
SETS = [('pitts30k_test', 'Pitts30k'), ('pitts250k_test', 'Pitts250k'), ('msls_val', 'MSLS'),
        ('sf_xl_small', 'SF-XL'), ('tokyo247', 'Tokyo 24/7'), ('amstertime', 'AmsterTime'),
        ('st_lucia', 'St Lucia'), ('eynsham', 'Eynsham'), ('nordland', 'Nordland'), ('svox', 'SVOX (all)'),
        ('svox_night', 'SVOX Night'), ('svox_overcast', 'SVOX Overcast'), ('svox_rain', 'SVOX Rain'),
        ('svox_snow', 'SVOX Snow'), ('svox_sun', 'SVOX Sun')]
KEYS = [k for k, _ in SETS]
NAME = dict(SETS)
GREY = RGBColor(0x5f, 0x5e, 0x58)


def load(suffix):
    d = pd.concat([pd.read_csv(f'results/seed{s}/test_results{suffix}.csv') for s in S])
    d = d[d.condition.isin(CONDS)]
    p = d.pivot_table(index=['seed', 'condition'], columns='dataset', values='R1')[KEYS]
    gain = pd.concat([p.loc[s] - p.loc[s].loc['none'] for s in S], keys=S)
    return dict(
        p=p,
        mean=d.groupby(['condition', 'dataset']).R1.mean().unstack()[KEYS].loc[CONDS],
        std=d.groupby(['condition', 'dataset']).R1.std().unstack()[KEYS].loc[CONDS],
        gain=gain.groupby(level=1).mean().loc[CONDS],
        up=gain.groupby(level=1).agg(lambda v: (v > 0).all()).loc[CONDS],
        dn=gain.groupby(level=1).agg(lambda v: (v < 0).all()).loc[CONDS],
    )


rgb, gray = load(''), load('_gray')
diff_runs = pd.concat([gray['p'].loc[s].loc[CONDS] - rgb['p'].loc[s].loc[CONDS] for s in S], keys=S)
diff = diff_runs.groupby(level=1).mean().loc[CONDS]
diff_up = diff_runs.groupby(level=1).agg(lambda v: (v > 0).all()).loc[CONDS]
diff_dn = diff_runs.groupby(level=1).agg(lambda v: (v < 0).all()).loc[CONDS]

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
doc.sections[0].left_margin = doc.sections[0].right_margin = Cm(2.0)
tno, fno = [0], [0]


def border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '000000')
    b.append(el)


def caption(text, kind='Table'):
    n = tno if kind == 'Table' else fno
    n[0] += 1
    p = doc.add_paragraph()
    r = p.add_run(f'{kind} {n[0]}. '); r.bold = True
    p.add_run(text)
    if kind == 'Table':
        p.paragraph_format.keep_with_next = True


def table(header, rows, widths, font=9, rules=(), bold=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        border(c, 'top', 12); border(c, 'bottom', 6)
    for i, cells in enumerate(rows, 1):
        rc = t.rows[i].cells
        for j, cell in enumerate(cells):
            parts = cell if isinstance(cell, list) else [(cell, False)]
            par = rc[j].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.keep_with_next = i < len(rows)
            for text, small in parts:
                r = par.add_run(text); r.font.size = Pt(font - 1.5 if small else font)
                r.bold = (i, j) in bold and not small
                if small:
                    r.font.color.rgb = GREY
        if cells[0] in rules:
            for c in rc: border(c, 'bottom', 4)
        if i == len(rows):
            for c in rc: border(c, 'bottom', 12)
    doc.add_paragraph()


def bullet(head, body):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(head + ' '); r.bold = True
    p.add_run(body)


def figure(path, width, text):
    doc.add_picture(path, width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(text, kind='Figure')


def fmt(v, digits=2):
    return '0.00' if round(v, digits) == 0 else f'{v:+.{digits}f}'


# ---------------------------------------------------------------- numbers used in the text
avg_r = rgb['gain'].mean(axis=1); avg_g = gray['gain'].mean(axis=1)
up_r = rgb['up'].sum(axis=1); up_g = gray['up'].sum(axis=1)
dn_r = rgb['dn'].sum(axis=1); dn_g = gray['dn'].sum(axis=1)
none_diff = diff.loc['none']
aug_night = diff.loc[SINGLE + MIXED, 'svox_night']
cost_sets = ['pitts30k_test', 'pitts250k_test', 'msls_val', 'st_lucia', 'svox']
avg_abs_r = rgb['mean'].mean(axis=1); avg_abs_g = gray['mean'].mean(axis=1)
ra = pd.concat([pd.read_csv(f'results/seed{s}/test_results.csv') for s in S])
ra_night = ra[(ra.condition == 'randaugment') & (ra.dataset == 'svox_night')].R1.mean()

# ---------------------------------------------------------------- title + setup
doc.add_heading('Grayscale results and comparison with RGB', level=1)
doc.add_paragraph('MixVPR (ResNet-50) trained on GSV-Cities for 50 epochs, three seeds per condition (190223, 1, 2), '
                  'tested on 15 test sets. Grayscale models see grayscale images (converted to three identical '
                  'channels) in training, validation and testing; RGB models see colour images throughout. All '
                  'other settings are identical. All numbers are Recall@1 (%).')

doc.add_heading('Key findings', level=2)
bullet('Removing colour has a small cost on most benchmarks.',
       f'Without augmentation, grayscale is {abs(none_diff[cost_sets].max()):.1f}–{abs(none_diff[cost_sets].min()):.1f} '
       f'points below RGB on Pitts30k, Pitts250k, MSLS, St Lucia and SVOX (all), and about 3 points below on '
       f'SF-XL ({none_diff["sf_xl_small"]:+.1f}), Tokyo 24/7 ({none_diff["tokyo247"]:+.1f}) and Nordland '
       f'({none_diff["nordland"]:+.1f}), lower in all three seeds.')
bullet('Grayscale is dramatically better at night.',
       f'On SVOX Night the grayscale baseline scores {gray["mean"].loc["none", "svox_night"]:.1f} % against '
       f'{rgb["mean"].loc["none", "svox_night"]:.1f} % for RGB ({none_diff["svox_night"]:+.1f} points), and with '
       f'augmentation the gap grows to {aug_night.min():+.1f} to {aug_night.max():+.1f} points. The best grayscale '
       f'model (geo_k1, {gray["mean"].loc["geo_k1", "svox_night"]:.1f} %) approaches RGB RandAugment '
       f'({ra_night:.1f} %), which includes brightness and colour changes. Both point to colour and lighting '
       'appearance, not geometry, as what makes night images hard for the RGB models.')
bullet('Grayscale is also better on AmsterTime and Eynsham.',
       f'AmsterTime: {diff.loc[CONDS, "amstertime"].min():+.1f} to {diff.loc[CONDS, "amstertime"].max():+.1f} points '
       '(its historical query photos are largely black-and-white, so grayscale training matches them better). '
       f'Eynsham, whose images are grayscale: {diff.loc[CONDS, "eynsham"].min():+.1f} to '
       f'{diff.loc[CONDS, "eynsham"].max():+.1f}.')
bullet('Geometric augmentation helps much more on grayscale.',
       f'Every augmentation improves the grayscale baseline in all three seeds on {int(up_g[SINGLE + MIXED].min())}–'
       f'{int(up_g[SINGLE + MIXED].max())} of 15 test sets and is never consistently worse; on RGB the range was '
       f'{int(up_r[SINGLE + MIXED].min())}–{int(up_r[SINGLE + MIXED].max())}, with some consistent losses. Average '
       f'gains are about twice as large ({avg_g[SINGLE + MIXED].min():+.2f} to {avg_g[SINGLE + MIXED].max():+.2f} vs. '
       f'{avg_r[SINGLE + MIXED].min():+.2f} to {avg_r[SINGLE + MIXED].max():+.2f}). Part of this is because the grayscale '
       'baseline starts lower on several test sets, which may leave more room to improve.')
bullet('The ranking of augmentations is similar.',
       f'geo_k2 again has the largest average gain ({avg_g["geo_k2"]:+.2f}) and geo_k1 is consistent on all '
       f'{int(up_g["geo_k1"])} test sets. Crop, the weakest augmentation on RGB ({avg_r["crop"]:+.2f}), is useful '
       f'on grayscale ({avg_g["crop"]:+.2f}), including Tokyo 24/7, where it hurt RGB '
       f'({rgb["gain"].loc["crop", "tokyo247"]:+.1f}) but helps grayscale ({gray["gain"].loc["crop", "tokyo247"]:+.1f}).')
n_rgb_ahead = int((diff.loc['geo_k2'] < 0).sum())
bullet('Which representation is better overall depends on the test sets.',
       f'Averaged over all 15 test sets, the best grayscale model ({avg_abs_g.idxmax()}, {avg_abs_g.max():.2f} %) is '
       f'ahead of the best RGB model (geo_k2, {avg_abs_r["geo_k2"]:.2f} %), but almost entirely because of SVOX Night '
       f'and AmsterTime. Comparing geo_k2 models test set by test set, RGB is ahead on {n_rgb_ahead} of 15.')

# ---------------------------------------------------------------- overview table
doc.add_heading('Overview', level=2)
caption('Augmentation effects on RGB and grayscale over all 15 test sets. Each representation is compared with '
        'its own no-augmentation baseline. "Better/worse" = test sets where the change had that direction in all '
        'three seeds; "average" = mean Recall@1 change over the 15 test sets.')
rows = []
for c in SINGLE + MIXED:
    rows.append([c, f'{up_r[c]}', f'{dn_r[c]}', fmt(avg_r[c]), f'{up_g[c]}', f'{dn_g[c]}', fmt(avg_g[c])])
table(['Condition', 'RGB better', 'RGB worse', 'RGB average', 'Gray better', 'Gray worse', 'Gray average'],
      rows, [2.6, 2.1, 2.1, 2.3, 2.1, 2.1, 2.3], font=9.5, rules=('shear',))

# ---------------------------------------------------------------- figures
doc.add_heading('Augmentation effects on grayscale', level=2)
figure('results/figures/gain_heatmap_gray_orangeblue_tall.png', 12.5,
       'Change in Recall@1 relative to the grayscale no-augmentation baseline, per test set (rows) and augmentation '
       '(columns); mean over three seeds. Orange: better; blue: worse; colour capped at ±5 (same scale as the RGB '
       'heatmap). Outlined, bold: same direction in all three seeds.')

doc.add_heading('Grayscale compared with RGB', level=2)
figure('results/figures/gray_minus_rgb_orangeblue.png', 13.0,
       'Grayscale minus RGB Recall@1 for the same augmentation condition, per test set (rows) and condition '
       '(columns, including no augmentation); mean over three seeds. Orange: grayscale better; blue: RGB better; '
       'colour capped at ±8. Outlined, bold: same direction for all three seeds.')
for h, b in [
    ('Blue rows (RGB better):', 'Pittsburgh, MSLS, SF-XL, St Lucia, SVOX (all) and Nordland, i.e. colour helps on '
     'ordinary daytime street scenes and on the summer/winter Nordland task.'),
    ('Orange rows (grayscale better):', 'SVOX Night (by far), AmsterTime and Eynsham for every condition; SVOX Rain '
     'and SVOX Sun once augmentation is used.'),
    ('Mixed rows:', 'Tokyo 24/7, SVOX Overcast and SVOX Snow depend on the augmentation.'),
]:
    bullet(h, b)

# ---------------------------------------------------------------- tables (landscape)
for part, (keys_part, label) in enumerate([(KEYS[:8], 'general benchmarks'), (KEYS[8:], 'seasonal and weather/lighting test sets')]):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    if sec.page_width < sec.page_height:
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Cm(1.8)
    doc.add_heading(f'Grayscale Recall@1 on the {label}', level=2)
    caption(f'Grayscale Recall@1 (%), mean ± standard deviation over three seeds, on the {label}. Small grey number: '
            'difference from the RGB model with the same condition.')
    rows, bold = [], set()
    for i, c in enumerate(CONDS, 1):
        cells = [c]
        for k in keys_part:
            cells.append([(f'{gray["mean"].loc[c, k]:.2f}', False), (f' ±{gray["std"].loc[c, k]:.2f}', True),
                          (f' ({fmt(diff.loc[c, k], 1)})', True)])
        rows.append(cells)
    for j, k in enumerate(keys_part, 1):
        best = gray['mean'][k].max()
        for i, c in enumerate(CONDS, 1):
            if round(gray['mean'].loc[c, k], 2) == round(best, 2):
                bold.add((i, j))
    w = [2.3] + [(25.4 - 2.3) / len(keys_part)] * len(keys_part)
    table(['Condition'] + [NAME[k] for k in keys_part], rows, w, font=8.5, rules=('none', 'shear'), bold=bold)

doc.save('results/grayscale_results.docx')
print('saved results/grayscale_results.docx; best gray avg', avg_abs_g.idxmax(), round(avg_abs_g.max(), 2),
      'best rgb geo_k2', round(avg_abs_r['geo_k2'], 2), 'rgb max', avg_abs_r.idxmax(), round(avg_abs_r.max(), 2))
