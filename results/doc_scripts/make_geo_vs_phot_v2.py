"""Geometric vs. photometric augmentation on RGB: comparison document. Run from MixVPR with python-docx."""
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

S = [190223, 1, 2]
GEO = ['crop', 'perspective', 'rotate', 'translate', 'shear']
MIX = ['geo_k1', 'geo_k2', 'geo_k3']
PHO = ['brightness', 'contrast', 'autocontrast', 'blur']
PMIX = ['phot_k1', 'phot_k2', 'phot_k3']
REF = ['randaugment']
ALLC = GEO + MIX + PHO + PMIX + REF
GEN = [('pitts30k_test', 'Pitts30k'), ('pitts250k_test', 'Pitts250k'), ('msls_val', 'MSLS'),
       ('sf_xl_small', 'SF-XL'), ('tokyo247', 'Tokyo 24/7'), ('amstertime', 'AmsterTime'),
       ('st_lucia', 'St Lucia'), ('eynsham', 'Eynsham')]
SEA = [('nordland', 'Nordland'), ('svox', 'SVOX (all)'), ('svox_night', 'SVOX Night'),
       ('svox_overcast', 'SVOX Overcast'), ('svox_rain', 'SVOX Rain'), ('svox_snow', 'SVOX Snow'),
       ('svox_sun', 'SVOX Sun')]
SETS = GEN + SEA
KEYS = [k for k, _ in SETS]
K_GEN, K_SEA = [k for k, _ in GEN], [k for k, _ in SEA]
NAME = dict(SETS)
GREY = RGBColor(0x5f, 0x5e, 0x58)

d = pd.concat([pd.read_csv(f'results/seed{s}/test_results.csv') for s in S])
p = d.pivot_table(index=['seed', 'condition'], columns='dataset', values='R1')[KEYS]
g = pd.concat([p.loc[s] - p.loc[s].loc['none'] for s in S], keys=S)
gm = g.groupby(level=1).mean()
gpos = g.groupby(level=1).agg(lambda v: int((v > 0).sum()))
up = g.groupby(level=1).agg(lambda v: bool((v > 0).all()))
dn = g.groupby(level=1).agg(lambda v: bool((v < 0).all()))
avg = gm.mean(axis=1)

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
doc.sections[0].left_margin = doc.sections[0].right_margin = Cm(2.0)
tno = [0]


def border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '000000')
    b.append(el)


def caption(text):
    tno[0] += 1
    p_ = doc.add_paragraph()
    r = p_.add_run(f'Table {tno[0]}. '); r.bold = True
    p_.add_run(text)
    p_.paragraph_format.keep_with_next = True


def table(header, rows, widths, font=9, rules=(), bold=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        border(c, 'top', 12); border(c, 'bottom', 6)
    for i, cells in enumerate(rows, 1):
        rc = t.rows[i].cells
        for j, cell in enumerate(cells):
            parts = cell if isinstance(cell, list) else [(cell, False)]
            par = rc[j].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.keep_with_next = i < len(rows)
            for text, small in parts:
                r = par.add_run(text); r.font.size = Pt(font - 1.5 if small else font)
                r.bold = (i, j) in bold and not small
                if small:
                    r.font.color.rgb = GREY
        if cells[0] in rules:
            for c in rc: border(c, 'bottom', 4)
        if i == len(rows):
            for c in rc: border(c, 'bottom', 12)
    doc.add_paragraph()


def section(landscape):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    w, h = s.page_width, s.page_height
    if (landscape and w < h) or (not landscape and w > h):
        s.page_width, s.page_height = h, w
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.left_margin = s.right_margin = Cm(1.8 if landscape else 2.0)


def f(v):
    return '0.00' if round(v, 2) == 0 else f'{v:+.2f}'


def label(c):
    return {'randaugment': 'RandAugment (ref.)'}.get(c, c)


# ---------------------------------------------------------------- intro
doc.add_heading('Geometric vs. photometric augmentation (RGB, three seeds)', level=1)
doc.add_paragraph('MixVPR (ResNet-50), RGB input, trained on GSV-Cities for 50 epochs with three seeds per condition; '
                  'tested on 15 test sets. All values are changes in Recall@1 (percentage points) from the same '
                  'no-augmentation model, computed within each seed and averaged.')
for t in [
    'Geometric: crop, perspective, rotate, translate, shear (single transforms) and geo_k1/k2/k3 (k random geometric '
    'transforms per image).',
    'Photometric: brightness, contrast, autocontrast, Gaussian blur (single transforms) and phot_k1/k2/k3 (k random '
    'photometric transforms per image).',
    'RandAugment is shown only as a reference: it mixes geometric and colour operations, so it belongs to neither family.',
    '"Consistent" = the change had the same direction in all three seeds.',
]:
    doc.add_paragraph(t, style='List Bullet')

# ---------------------------------------------------------------- overview
doc.add_heading('1. Overview over all 15 test sets', level=2)
caption('Each condition over all 15 test sets: test sets improved / worsened in all three seeds, and the average '
        'R@1 change.')
rows = [[label(c), f'{int(up.loc[c].sum())}', f'{int(dn.loc[c].sum())}', f(avg[c])] for c in ALLC]
table(['Condition', 'Better (3/3)', 'Worse (3/3)', 'Average change'], rows, [4.0, 3.0, 3.0, 3.4], font=10,
      rules=('shear', 'geo_k3', 'blur', 'phot_k3'))

doc.add_picture('results/figures/family_bars_rgb.png', width=Cm(17))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
pf = doc.add_paragraph()
r = pf.add_run('Figure 1. '); r.bold = True
pf.add_run('Average R@1 change from no augmentation for each condition on the general benchmarks, the seasonal and '
           'weather test sets (without SVOX Night) and SVOX Night, coloured by family. Error bars: standard deviation '
           'over the three seeds. Dashed line: RandAugment reference.')

# ---------------------------------------------------------------- family averages
doc.add_heading('2. Family averages', level=2)
caption('Average R@1 change of each family\'s single transforms (five geometric, four photometric) and of its best '
        'single transform, per test-set group.')
rows = []
for fam, conds in [('Geometric (single, mean)', GEO), ('Photometric (single, mean)', PHO)]:
    rows.append([fam] + [f(gm.loc[conds, ks].mean(axis=1).mean()) for ks in (K_GEN, K_SEA, KEYS)])
for fam, conds in [('Best geometric single', GEO), ('Best photometric single', PHO)]:
    cells = [fam]
    for ks in (K_GEN, K_SEA, KEYS):
        a = gm.loc[conds, ks].mean(axis=1)
        cells.append([(f(a.max()), False), (f' ({a.idxmax()})', True)])
    rows.append(cells)
for fam, conds in [('Best geometric mixed', MIX), ('Best photometric mixed', PMIX)]:
    cells = [fam]
    for ks in (K_GEN, K_SEA, KEYS):
        a = gm.loc[conds, ks].mean(axis=1)
        cells.append([(f(a.max()), False), (f' ({a.idxmax()})', True)])
    rows.append(cells)
rows.append(['RandAugment (ref.)'] + [f(gm.loc['randaugment', ks].mean()) for ks in (K_GEN, K_SEA, KEYS)])
table(['', 'General (8)', 'Seasonal/weather (7)', 'All (15)'], rows, [5.0, 3.6, 3.8, 3.4], font=10,
      rules=('Photometric (single, mean)', 'Best photometric single', 'Best photometric mixed'))

# ---------------------------------------------------------------- per test set winner
doc.add_heading('3. Which family helps most on each test set', level=2)
caption('Per test set: best single geometric and best single photometric transform (by mean R@1 change), the best '
        'mixed condition of each family (geo_k2, phot_k3) and the RandAugment reference. Bold: consistent in all three seeds.')
rows, bold = [], set()
for i, (k, name) in enumerate(SETS, 1):
    bg, bp = gm.loc[GEO, k].idxmax(), gm.loc[PHO, k].idxmax()
    cells = [name]
    for j, c in enumerate([bg, bp, 'geo_k2', 'phot_k3', 'randaugment'], 1):
        txt = [(f(gm.loc[c, k]), False)] + ([(f' ({c})', True)] if j <= 2 else [])
        cells.append(txt)
        if up.loc[c, k] or dn.loc[c, k]:
            bold.add((i, j))
    rows.append(cells)
table(['Test set', 'Best geo. single', 'Best phot. single', 'geo_k2', 'phot_k3', 'RandAugment'], rows,
      [2.8, 3.4, 3.4, 1.8, 1.8, 2.2], font=8.5, rules=('Eynsham',), bold=bold)

# ---------------------------------------------------------------- full gain tables
for group, gname in [(GEN, 'general benchmarks'), (SEA, 'seasonal and weather/lighting test sets')]:
    section(True)
    doc.add_heading(f'4. All conditions on the {gname}', level=2)
    caption(f'R@1 change from no augmentation on the {gname}, mean over three seeds; brackets: seeds with a positive '
            'change. Bold: consistent in all three seeds. Rows: geometric single, geometric mixed, photometric, reference.')
    keys = [k for k, _ in group]
    rows, bold = [], set()
    for i, c in enumerate(ALLC, 1):
        rows.append([label(c)] + [[(f(gm.loc[c, k]), False), (f' ({gpos.loc[c, k]}/3)', True)] for k in keys])
        for j, k in enumerate(keys, 1):
            if up.loc[c, k] or dn.loc[c, k]:
                bold.add((i, j))
    table(['Condition'] + [s for _, s in group], rows, [3.6] + [(25.8 - 3.6) / len(keys)] * len(keys), font=8.5,
          rules=('shear', 'geo_k3', 'blur', 'phot_k3'), bold=bold)

# ---------------------------------------------------------------- key findings
section(False)
doc.add_heading('5. Key findings', level=2)
fam_g = gm.loc[GEO].mean(axis=1).mean(); fam_p = gm.loc[PHO].mean(axis=1).mean()
night = {c: (gm.loc[c, 'svox_night'], bool(up.loc[c, 'svox_night'])) for c in GEO + MIX + PHO}
nights_g = [gm.loc[c, 'svox_night'] for c in GEO + MIX]
cn = gm.loc['contrast']
pavg = {c: avg[c] for c in ALLC}
non_ra = GEO + MIX + PHO + PMIX
beats_ra = [(k, [c for c in non_ra if round(gm.loc[c, k], 2) >= round(gm.loc['randaugment', k], 2)]) for k in KEYS]
beats_ra = [(k, cs) for k, cs in beats_ra if cs]
pk = gm.loc['phot_k3']
findings = [
    ('Single transforms: the two families help about equally on average.',
     f'The five geometric transforms average {fam_g:+.2f} and the four photometric transforms {fam_p:+.2f} over the 15 '
     f'test sets; the best of each is similar (contrast {avg["contrast"]:+.2f}, perspective {avg["perspective"]:+.2f}).'),
    ('They help in different places.',
     f'Geometric augmentation gives the larger gains on the urban benchmarks (SF-XL: geo_k2 '
     f'{gm.loc["geo_k2", "sf_xl_small"]:+.1f}; MSLS: perspective {gm.loc["perspective", "msls_val"]:+.1f}) and on '
     f'Nordland among single transforms (perspective {gm.loc["perspective", "nordland"]:+.1f}). Photometric augmentation '
     f'dominates on lighting change: SVOX Night (contrast {cn["svox_night"]:+.1f}, phot_k3 {pk["svox_night"]:+.1f}, '
     f'against at most {max(nights_g):+.1f} for any geometric condition).'),
    ('Combining photometric transforms gives the strongest non-reference condition.',
     f'phot_k3 has the largest average gain of all geometric and photometric conditions ({avg["phot_k3"]:+.2f}; geo_k2 '
     f'{avg["geo_k2"]:+.2f}), improves {int(up.loc["phot_k3"].sum())} of 15 test sets in all three seeds and is never '
     f'consistently worse. Its gains are largest under lighting and seasonal change: SVOX Night {pk["svox_night"]:+.1f}, '
     f'SVOX Sun {pk["svox_sun"]:+.1f}, Nordland {pk["nordland"]:+.1f}, Tokyo 24/7 {pk["tokyo247"]:+.1f}.'),
    ('The number of combined photometric transforms does not act as a simple dial.',
     f'phot_k1 ({avg["phot_k1"]:+.2f}, {int(up.loc["phot_k1"].sum())} test sets) is ahead of phot_k2 '
     f'({avg["phot_k2"]:+.2f}, {int(up.loc["phot_k2"].sum())} test sets), and phot_k3 is clearly best; phot_k2 also has '
     'the largest seed-to-seed variation. With three seeds this ordering should be read with caution.'),
    ('Contrast is the most useful single photometric transform; autocontrast and blur are weak.',
     f'Contrast ({avg["contrast"]:+.2f}) has a night gain of {cn["svox_night"]:+.1f} that is consistent in all three '
     f'seeds; brightness is reliable but modest ({avg["brightness"]:+.2f}); autocontrast ({avg["autocontrast"]:+.2f}) '
     f'has almost no effect and blur ({avg["blur"]:+.2f}) helps mainly on Nordland.'),
    ('RandAugment is still ahead on average, but not everywhere.',
     f'RandAugment (reference) averages {avg["randaugment"]:+.2f}, mainly through SVOX Night '
     f'({gm.loc["randaugment", "svox_night"]:+.1f}) and SVOX Sun ({gm.loc["randaugment", "svox_sun"]:+.1f}). On '
     f'{len(beats_ra)} test sets a geometric or photometric condition matches or beats it: '
     + '; '.join((f'{NAME[k]} (' + ', '.join(f'{c} {gm.loc[c, k]:+.2f}' for c in cs)
                  + f' vs. {gm.loc["randaugment", k]:+.2f})') if len(cs) <= 3 else
                 (f'{NAME[k]} ({len(cs)} of {len(non_ra)} conditions, best {gm.loc[cs, k].idxmax()} '
                  f'{gm.loc[cs, k].max():+.2f}, vs. {gm.loc["randaugment", k]:+.2f}; all changes under one point)')
                 for k, cs in beats_ra)
     + '. Mixing photometric transforms closes part of the night gap (phot_k3 '
     f'{pk["svox_night"]:+.1f} vs. {gm.loc["randaugment", "svox_night"]:+.1f}) but not most of it.'),
    ('Near-ceiling test sets are unaffected.',
     'St Lucia and SVOX (all) change by less than half a point for every geometric and photometric condition.'),
]
for head, body in findings:
    p_ = doc.add_paragraph(style='List Bullet')
    r = p_.add_run(head + ' '); r.bold = True
    p_.add_run(body)

doc.save('results/geometric_vs_photometric_rgb.docx')
print('saved', tno[0], 'tables | contrast night consistent:', bool(up.loc['contrast', 'svox_night']),
      '| contrast night per seed:', g.xs('contrast', level=1)['svox_night'].round(2).tolist())
