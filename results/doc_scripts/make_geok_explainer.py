"""Plain-language explainer for geo_k1/k2/k3. Run from the MixVPR folder with python-docx on PYTHONPATH."""
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SEEDS = [190223, 1, 2]
SINGLE = ['crop', 'perspective', 'rotate', 'translate', 'shear']
MIXED = ['geo_k1', 'geo_k2', 'geo_k3']
HOURS = {'geo_k1': 6.4, 'geo_k2': 7.7, 'geo_k3': 9.3, 'single': 6.5}  # mean wall-clock per training run
SETS = [('pitts30k_test', 'Pitts30k'), ('pitts250k_test', 'Pitts250k'), ('msls_val', 'MSLS'),
        ('sf_xl_small', 'SF-XL'), ('tokyo247', 'Tokyo 24/7'), ('amstertime', 'AmsterTime'),
        ('st_lucia', 'St Lucia'), ('eynsham', 'Eynsham'), ('nordland', 'Nordland'), ('svox', 'SVOX (all)'),
        ('svox_night', 'SVOX Night'), ('svox_overcast', 'SVOX Overcast'), ('svox_rain', 'SVOX Rain'),
        ('svox_snow', 'SVOX Snow'), ('svox_sun', 'SVOX Sun')]
KEYS = [k for k, _ in SETS]
GREY = RGBColor(0x5f, 0x5e, 0x58)

d = pd.concat([pd.read_csv(f'results/seed{s}/test_results.csv') for s in SEEDS])
mean = d.groupby(['condition', 'dataset']).R1.mean().unstack()[KEYS]
p = d.pivot_table(index=['seed', 'condition'], columns='dataset', values='R1')[KEYS]
gain = pd.concat([p.loc[s] - p.loc[s].loc['none'] for s in SEEDS], keys=SEEDS)
up = gain.groupby(level=1).agg(lambda v: int((v > 0).sum())) == 3
down = gain.groupby(level=1).agg(lambda v: int((v < 0).sum())) == 3
avg_gain = gain.groupby(level=1).mean().mean(axis=1)

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
sec = doc.sections[0]; sec.left_margin = sec.right_margin = Cm(2.2)
tno, fno = [0], [0]


def border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '000000')
    b.append(el)


def caption(text, kind='Table'):
    n = tno if kind == 'Table' else fno
    n[0] += 1
    p_ = doc.add_paragraph()
    r = p_.add_run(f'{kind} {n[0]}. '); r.bold = True
    p_.add_run(text)
    if kind == 'Table':
        p_.paragraph_format.keep_with_next = True


def table(header, rows, widths, font=10, rules=(), bold_rows=()):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for row in tbl.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    for j, h in enumerate(header):
        c = tbl.rows[0].cells[j]
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        border(c, 'top', 12); border(c, 'bottom', 6)
    for i, cells in enumerate(rows, 1):
        rc = tbl.rows[i].cells
        for j, text in enumerate(cells):
            par = rc[j].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(text); r.font.size = Pt(font); r.bold = cells[0] in bold_rows
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.keep_with_next = i < len(rows)
        if cells[0] in rules:
            for c in rc: border(c, 'bottom', 4)
        if i == len(rows):
            for c in rc: border(c, 'bottom', 12)
    doc.add_paragraph()


def bullets(items):
    for t in items:
        doc.add_paragraph(t, style='List Bullet')


k2_vs_p = mean.loc['geo_k2'] - mean.loc['perspective']
best_single = mean.loc[SINGLE].max()
best_name = mean.loc[SINGLE].idxmax()
k2_vs_best = mean.loc['geo_k2'] - best_single
n_sets = len(KEYS)

# ---------------------------------------------------------------- title + summary
doc.add_heading('Mixed geometric augmentation: geo_k1, geo_k2 and geo_k3', level=1)
doc.add_paragraph('MixVPR (ResNet-50), RGB input, trained on GSV-Cities for 50 epochs with three seeds per '
                  'condition (190223, 1, 2), tested on 15 test sets. All numbers are Recall@1 (%): the share '
                  'of queries whose top-ranked database image shows the correct place.')

doc.add_heading('In short', level=2)
bullets([
    'geo_k1, geo_k2 and geo_k3 apply 1, 2 or 3 geometric transforms to every training image. Which '
    'transforms are used is chosen at random for each image.',
    f'geo_k2 is the best of the three and the best geometric condition overall. It beats no augmentation in '
    f'all three seeds on {int(up.loc["geo_k2"].sum())} of {n_sets} test sets and never makes things '
    f'consistently worse.',
    f'Compared with the best single transform (perspective), geo_k2 is only slightly better on average: '
    f'{k2_vs_p.mean():+.2f} points, better on {int((k2_vs_p > 0).sum())} of {n_sets} test sets. Its main '
    f'advantage is that it helps on more test sets, not that it scores much higher.',
    f'geo_k3 gives no extra benefit over geo_k2 and takes longer to train; geo_k1 is weaker than both.',
])

# ---------------------------------------------------------------- what they are
doc.add_heading('1. What geo_k1, geo_k2 and geo_k3 are', level=2)
doc.add_paragraph('The five single geometric augmentations are crop, perspective, rotate, translate and shear. '
                  'Each single condition applies one fixed transform to every training image. The mixed '
                  'conditions instead pick transforms at random from these five, separately for every image, '
                  'every time it is loaded:')
caption('What each mixed condition does.')
table(['Condition', 'Transforms per image', 'How they are chosen', 'Training time per run'],
      [['geo_k1', '1', '1 of the 5, at random', f'≈ {HOURS["geo_k1"]:.1f} h'],
       ['geo_k2', '2 (different)', '2 of the 5, at random, random order', f'≈ {HOURS["geo_k2"]:.1f} h'],
       ['geo_k3', '3 (different)', '3 of the 5, at random, random order', f'≈ {HOURS["geo_k3"]:.1f} h']],
      [2.6, 3.6, 6.2, 3.8])
bullets([
    'Because the choice is new for every image, over the whole training each model sees every transform '
    '(geo_k1), every pair (geo_k2) or every triple (geo_k3) many times. A mixed condition is therefore a '
    'rule ("apply two random transforms"), not one particular combination.',
    'The transforms use exactly the same settings as in the single conditions (e.g. rotation up to ±10°), so '
    'the only difference is how many are applied and that they vary.',
    'geo_k1 applies the same amount of augmentation as a single condition (one transform per image) but with '
    'variety. geo_k2 and geo_k3 apply more augmentation per image.',
])
doc.add_picture('results/figures/geo_k_examples.png', width=Cm(15.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption('One GSV-Cities training image (top) and three random examples of what each mixed condition can '
        'produce from it. The title of each example lists the transforms that were drawn, in the order they '
        'were applied. Black borders come from rotating, shearing or shifting the image.', kind='Figure')

# ---------------------------------------------------------------- compare with each other
doc.add_heading('2. How they compare with each other', level=2)
caption(f'Summary of the mixed conditions over all {n_sets} test sets. "Better/worse in all 3 seeds" counts test '
        'sets where the change from no augmentation had the same direction in every seed. Average change is '
        'the R@1 change from no augmentation, averaged over the test sets.')
table(['Condition', 'Better in all 3 seeds', 'Worse in all 3 seeds', 'Average change', 'Training time'],
      [[c, f'{int(up.loc[c].sum())} of {n_sets}', f'{int(down.loc[c].sum())} of {n_sets}',
        f'{avg_gain.loc[c]:+.2f}', f'≈ {HOURS[c]:.1f} h'] for c in MIXED],
      [2.6, 3.4, 3.4, 3.0, 3.0], bold_rows=('geo_k2',))
k2_vs_k1 = mean.loc['geo_k2'] - mean.loc['geo_k1']
k3_vs_k2 = mean.loc['geo_k3'] - mean.loc['geo_k2']
bullets([
    f'geo_k2 beats geo_k1 on {int((k2_vs_k1 > 0).sum())} of {n_sets} test sets, by {k2_vs_k1.mean():+.2f} points '
    f'on average. The biggest differences are on SVOX Night ({k2_vs_k1["svox_night"]:+.1f}), SVOX Sun '
    f'({k2_vs_k1["svox_sun"]:+.1f}), Tokyo 24/7 ({k2_vs_k1["tokyo247"]:+.1f}) and SF-XL '
    f'({k2_vs_k1["sf_xl_small"]:+.1f}). One random transform per image is not enough; two is clearly better.',
    f'geo_k3 and geo_k2 are almost the same: geo_k3 is ahead on {int((k3_vs_k2 > 0).sum())} test sets but '
    f'behind on average ({k3_vs_k2.mean():+.2f}), and every difference is small (largest '
    f'{k3_vs_k2.abs().max():.1f} points, on {dict(SETS)[k3_vs_k2.abs().idxmax()]}). A third transform adds '
    f'about 20 % training time for no gain.',
    'Conclusion: geo_k2 is the best choice of the three.',
])

# ---------------------------------------------------------------- compare with single
doc.add_heading('3. How they compare with the single augmentations', level=2)
caption(f'Mixed and single conditions over all {n_sets} test sets (same measures as Table 2).')
rows = [[c, f'{int(up.loc[c].sum())} of {n_sets}', f'{int(down.loc[c].sum())} of {n_sets}', f'{avg_gain.loc[c]:+.2f}']
        for c in SINGLE + MIXED]
table(['Condition', 'Better in all 3 seeds', 'Worse in all 3 seeds', 'Average change'], rows,
      [3.0, 3.8, 3.8, 3.6], rules=('shear',), bold_rows=('geo_k2', 'perspective'))
bullets([
    f'Measured by how many test sets improve reliably, geo_k2 and geo_k3 ({int(up.loc["geo_k2"].sum())} each) beat '
    f'every single transform; the best single transform, perspective, manages {int(up.loc["perspective"].sum())}. '
    'Neither geo_k2 nor geo_k3 is consistently worse than no augmentation anywhere.',
    f'Measured by average gain, geo_k2 ({avg_gain.loc["geo_k2"]:+.2f}) is ahead of perspective '
    f'({avg_gain.loc["perspective"]:+.2f}) by {avg_gain.loc["geo_k2"] - avg_gain.loc["perspective"]:.2f} points: '
    'a small margin.',
    f'geo_k1 ({avg_gain.loc["geo_k1"]:+.2f}) is worse than perspective and shear alone: one random transform '
    'per image does not beat simply always using a good transform.',
])

doc.add_paragraph('Table 4 shows the comparison test set by test set. "Best single" is whichever of the five '
                  'single transforms scored highest on that test set, so it is a demanding comparison: a '
                  'different single transform wins on different test sets, while geo_k2 is one fixed condition.')
caption('Recall@1 (%) per test set, mean over three seeds. The last two columns show how much geo_k2 is ahead of '
        '(+) or behind (−) perspective and the best single transform on that test set.')
rows = []
for k, name in SETS:
    rows.append([name, f'{mean.loc["none", k]:.2f}', f'{mean.loc["perspective", k]:.2f}',
                 f'{best_single[k]:.2f} ({best_name[k]})', f'{mean.loc["geo_k1", k]:.2f}',
                 f'{mean.loc["geo_k2", k]:.2f}', f'{mean.loc["geo_k3", k]:.2f}',
                 f'{k2_vs_p[k]:+.2f}', f'{k2_vs_best[k]:+.2f}'])
rows.append(['Average', f'{mean.loc["none"].mean():.2f}', f'{mean.loc["perspective"].mean():.2f}', '—',
             f'{mean.loc["geo_k1"].mean():.2f}', f'{mean.loc["geo_k2"].mean():.2f}',
             f'{mean.loc["geo_k3"].mean():.2f}', f'{k2_vs_p.mean():+.2f}', f'{k2_vs_best.mean():+.2f}'])
table(['Test set', 'none', 'persp.', 'Best single', 'geo_k1', 'geo_k2', 'geo_k3', 'k2 vs persp.',
       'k2 vs best single'], rows, [2.3, 1.4, 1.4, 3.2, 1.5, 1.5, 1.5, 1.7, 1.9], font=8.5,
      rules=('SVOX Sun',), bold_rows=('Average',))
big_up = k2_vs_p.sort_values(ascending=False)
bullets([
    f'Where geo_k2 is most ahead of perspective: SVOX Sun ({k2_vs_p["svox_sun"]:+.1f}) and SF-XL '
    f'({k2_vs_p["sf_xl_small"]:+.1f}). Elsewhere the differences are under one point and mostly within the '
    f'variation between seeds.',
    f'Where perspective is most ahead: Nordland, by {-k2_vs_p["nordland"]:.1f} points. On Nordland, always '
    f'applying perspective works better than mixing.',
    f'Against the best single transform on each test set, geo_k2 is level on average '
    f'({k2_vs_best.mean():+.2f}), ahead on {int((k2_vs_best > 0).sum())} of {n_sets}. In other words, one '
    'mixed condition does about as well as picking the ideal single transform separately for each test set, '
    'which you could not do in practice without knowing the test set in advance.',
])

# ---------------------------------------------------------------- caveats
doc.add_heading('4. Things to keep in mind', level=2)
bullets([
    'Most per-test-set differences in Table 4 are small (under one point) and within the variation between '
    'seeds. The reliable findings are the counts in Tables 2–3 (same direction in all three seeds) and the larger '
    'gaps such as SVOX Sun, SF-XL and Nordland.',
    'The results cannot say which specific combination of transforms is responsible, because every model saw '
    'all combinations. To test a particular pair (e.g. perspective + shear), it would have to be trained as its '
    'own condition.',
    'Night images are not solved by any geometric condition; that needs photometric (colour/brightness) '
    'augmentation, as the RandAugment reference results show.',
])

out = 'results/geo_k_explained.docx'
doc.save(out)
print('saved', out, 'tables', tno[0], 'figures', fno[0])
