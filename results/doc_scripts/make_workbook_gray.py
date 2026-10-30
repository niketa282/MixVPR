"""Data-only workbook (one table per page + notes box). Run from MixVPR with python-docx on PYTHONPATH."""
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SEEDS = [190223, 1, 2]
SINGLE = ['crop', 'perspective', 'rotate', 'translate', 'shear']
MIXED = ['geo_k1', 'geo_k2', 'geo_k3']
CONDS = ['none'] + SINGLE + MIXED
RULES = ('none', 'shear')
GEN = [('pitts30k_test', 'Pitts30k'), ('pitts250k_test', 'Pitts250k'), ('msls_val', 'MSLS'),
       ('sf_xl_small', 'SF-XL'), ('tokyo247', 'Tokyo 24/7'), ('amstertime', 'AmsterTime'),
       ('st_lucia', 'St Lucia'), ('eynsham', 'Eynsham')]
SEA = [('nordland', 'Nordland'), ('svox', 'SVOX (all)'), ('svox_night', 'SVOX Night'),
       ('svox_overcast', 'SVOX Overcast'), ('svox_rain', 'SVOX Rain'), ('svox_snow', 'SVOX Snow'),
       ('svox_sun', 'SVOX Sun')]
ABOUT = {'pitts30k_test': 'Pittsburgh street view (same city as validation)', 'pitts250k_test': 'larger Pittsburgh set',
         'msls_val': 'dashcams, several cities, mixed conditions', 'sf_xl_small': 'San Francisco, crowd-sourced query photos',
         'tokyo247': 'Tokyo, day / sunset / night queries', 'amstertime': 'historical vs. modern Amsterdam photos',
         'st_lucia': 'suburban driving, different times of day', 'eynsham': 'rural roads, grayscale camera',
         'nordland': 'railway, summer database vs. winter queries', 'svox': 'Oxford, all query conditions',
         'svox_night': 'Oxford, night queries', 'svox_overcast': 'Oxford, overcast queries',
         'svox_rain': 'Oxford, rain queries', 'svox_snow': 'Oxford, snow queries', 'svox_sun': 'Oxford, strong sun queries'}
SETS = GEN + SEA
KEYS = [k for k, _ in SETS]
GREY = RGBColor(0x5f, 0x5e, 0x58)

d = pd.concat([pd.read_csv(f'results/seed{s}/test_results_gray.csv') for s in SEEDS])
drgb = pd.concat([pd.read_csv(f'results/seed{s}/test_results.csv') for s in SEEDS])
drgb = drgb[drgb.condition.isin(CONDS)]
prgb = drgb.pivot_table(index=['seed', 'condition'], columns='dataset', values='R1')[KEYS]
mean_rgb = drgb.groupby(['condition', 'dataset']).R1.mean().unstack()[KEYS]
d = d[d.condition.isin(CONDS)]
mean = d.groupby(['condition', 'dataset']).R1.mean().unstack()[KEYS]
std = d.groupby(['condition', 'dataset']).R1.std().unstack()[KEYS]
p = d.pivot_table(index=['seed', 'condition'], columns='dataset', values='R1')[KEYS]
gain = pd.concat([p.loc[s] - p.loc[s].loc['none'] for s in SEEDS], keys=SEEDS)
gmean = gain.groupby(level=1).mean()
npos = gain.groupby(level=1).agg(lambda v: int((v > 0).sum()))
nneg = gain.groupby(level=1).agg(lambda v: int((v < 0).sum()))
counts = d.groupby('dataset')[['num_db', 'num_queries']].first()

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
doc.sections[0].left_margin = doc.sections[0].right_margin = Cm(2.0)
pno = [0]


def border(cell, edge, sz, val='single'):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), val); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '7f7f7f' if val != 'single' or sz < 6 else '000000')
    b.append(el)


def new_page(landscape):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    w, h = s.page_width, s.page_height
    if (landscape and w < h) or (not landscape and w > h):
        s.page_width, s.page_height = h, w
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.left_margin = s.right_margin = Cm(1.8 if landscape else 2.0)
    s.top_margin = s.bottom_margin = Cm(1.6)


def heading(text, sub=None):
    pno[0] += 1
    doc.add_heading(f'Page {pno[0]}. {text}', level=2)
    if sub:
        r = doc.add_paragraph().add_run(sub); r.font.size = Pt(9.5); r.font.color.rgb = GREY


def table(header, rows, widths, font=9.5, bold_cells=()):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for row in tbl.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    for j, h in enumerate(header):
        c = tbl.rows[0].cells[j]
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        border(c, 'top', 12); border(c, 'bottom', 6)
    for i, cells in enumerate(rows, 1):
        rc = tbl.rows[i].cells
        for j, cell in enumerate(cells):
            parts = cell if isinstance(cell, list) else [(cell, False)]
            par = rc[j].paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            for text, small in parts:
                r = par.add_run(text); r.font.size = Pt(font - 1.5 if small else font)
                r.bold = (i, j) in bold_cells and not small
                if small:
                    r.font.color.rgb = GREY
        if cells[0] in RULES:
            for c in rc: border(c, 'bottom', 4)
        if i == len(rows):
            for c in rc: border(c, 'bottom', 12)


def notes(height_cm, width_cm):
    doc.add_paragraph()
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    row = tbl.rows[0]
    row.height = Cm(height_cm); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    c = row.cells[0]; c.width = Cm(width_cm)
    for e in ('top', 'bottom', 'left', 'right'):
        border(c, e, 6, val='dashed')
    r = c.paragraphs[0].add_run('My thoughts:'); r.bold = True; r.font.size = Pt(10)
    c.add_paragraph('')


def best_bold(row_vals, start_row=1, col=None):
    """Return cells (row, col) holding the max per column for the given dict col->list of values."""
    out = set()
    for j, vals in row_vals.items():
        m = max(vals)
        for i, v in enumerate(vals):
            if round(v, 2) == round(m, 2):
                out.add((i + start_row, j))
    return out


# ------------------------------------------------------------------ cover
doc.add_heading('Grayscale: single vs. mixed geometric augmentation, data workbook', level=1)
doc.add_paragraph('One table per page, with a box underneath for your own notes. Type straight into the boxes, '
                  'save, and send the file back.')
doc.add_heading('What the numbers are', level=3)
for t in [
    'All models here were trained, validated and tested on grayscale images (three identical channels). '
    'Pages 21 onwards compare them with the RGB models.',
    'All values are Recall@1 (%): the share of queries whose top-ranked database image shows the right place. '
    'Higher is better.',
    '"mean ± std": average over the three training seeds (190223, 1, 2) ± how much it varied between them.',
    '"change": difference from no augmentation (none), worked out per seed and then averaged. The small '
    'number in brackets, e.g. (3/3), is in how many of the three seeds the change was positive.',
    'Bold: the highest value in that column.',
    'Rows are grouped: none | the five single transforms | the three mixed conditions.',
]:
    doc.add_paragraph(t, style='List Bullet')
doc.add_heading('The conditions', level=3)
table(['Condition', 'What happens to each training image'],
      [['none', 'nothing'],
       ['crop', 'random zoom-in crop (70–100 % of the image)'],
       ['perspective', 'random perspective tilt (strength 0.2)'],
       ['rotate', 'random rotation, up to ±10°'],
       ['translate', 'random shift, up to 5 % of width/height'],
       ['shear', 'random shear, up to ±5° in x and y'],
       ['geo_k1', '1 of the five transforms above, chosen at random per image'],
       ['geo_k2', '2 different ones, chosen at random per image, random order'],
       ['geo_k3', '3 different ones, chosen at random per image, random order']],
      [3.0, 12.0], font=10)
notes(4, 15)

# ------------------------------------------------------------------ overview
new_page(False)
heading('Overview over all 15 test sets',
        'Better / worse in 3/3 seeds = number of test sets where the change from none had that direction in every '
        'seed. Average change = mean R@1 change from none over the 15 test sets.')
rows, vals = [], {1: [], 2: [], 3: []}
for c in CONDS[1:]:
    up_, dn_, av = int((npos.loc[c] == 3).sum()), int((nneg.loc[c] == 3).sum()), gmean.loc[c].mean()
    rows.append([c, f'{up_} of 15', f'{dn_} of 15', f'{av:+.2f}'])
    vals[1].append(up_); vals[3].append(av)
bc = best_bold({1: vals[1], 3: vals[3]})
table(['Condition', 'Better in 3/3 seeds', 'Worse in 3/3 seeds', 'Average change'], rows, [3.4, 4.0, 4.0, 4.0],
      font=10.5, bold_cells=bc)
notes(9, 15.4)

# ------------------------------------------------------------------ R@1 tables (landscape)
for group, gname in [(GEN, 'general benchmarks'), (SEA, 'seasonal and weather/lighting change')]:
    new_page(True)
    heading(f'Recall@1 on the {gname}', 'mean ± std over three seeds')
    hdr = ['Condition'] + [s for _, s in group]
    rows = [[c] + [[(f'{mean.loc[c, k]:.2f}', False), (f' ± {std.loc[c, k]:.2f}', True)] for k, _ in group]
            for c in CONDS]
    bc = best_bold({j: [mean.loc[c, k] for c in CONDS] for j, (k, _) in enumerate(group, 1)})
    table(hdr, rows, [2.4] + [(26.0 - 2.4) / len(group)] * len(group), font=9, bold_cells=bc)
    notes(4.5, 26.0)

    new_page(True)
    heading(f'Change from no augmentation on the {gname}',
            'change in R@1 (percentage points), mean over seeds; brackets: seeds with a positive change')
    rows = [[c] + [[(f'{gmean.loc[c, k]:+.2f}', False), (f' ({npos.loc[c, k]}/3)', True)] for k, _ in group]
            for c in CONDS[1:]]
    bc = best_bold({j: [gmean.loc[c, k] for c in CONDS[1:]] for j, (k, _) in enumerate(group, 1)})
    table(hdr, rows, [2.4] + [(26.0 - 2.4) / len(group)] * len(group), font=9, bold_cells=bc)
    notes(5.5, 26.0)

# ------------------------------------------------------------------ one page per test set
for k, name in SETS:
    new_page(False)
    heading(f'{name}', f'{ABOUT[k]} · {counts.loc[k, "num_db"]:,} database images, '
                       f'{counts.loc[k, "num_queries"]:,} queries')
    rows = []
    for c in CONDS:
        seeds = [f'{p.loc[(s, c), k]:.2f}' for s in SEEDS]
        ch = '—' if c == 'none' else [(f'{gmean.loc[c, k]:+.2f}', False), (f' ({npos.loc[c, k]}/3)', True)]
        rows.append([c, [(f'{mean.loc[c, k]:.2f}', False), (f' ± {std.loc[c, k]:.2f}', True)], ch] + seeds)
    bc = best_bold({1: [mean.loc[c, k] for c in CONDS], 2: [-1e9] + [gmean.loc[c, k] for c in CONDS[1:]]})
    table(['Condition', 'R@1 mean ± std', 'Change from none', 'Seed 190223', 'Seed 1', 'Seed 2'], rows,
          [2.6, 3.4, 3.0, 2.2, 2.0, 2.0], font=10, bold_cells=bc)
    notes(8, 15.2)


# ------------------------------------------------------------------ grayscale vs RGB
diff_runs = pd.concat([p.loc[s].loc[CONDS] - prgb.loc[s].loc[CONDS] for s in SEEDS], keys=SEEDS)
dmean = diff_runs.groupby(level=1).mean()
dpos = diff_runs.groupby(level=1).agg(lambda v: int((v > 0).sum()))
for group, gname in [(GEN, 'general benchmarks'), (SEA, 'seasonal and weather/lighting change')]:
    new_page(True)
    heading(f'Grayscale vs. RGB on the {gname}',
            'grayscale minus RGB Recall@1 for the same condition (percentage points), mean over seeds; '
            'brackets: seeds where grayscale was higher. Positive = grayscale better.')
    hdr = ['Condition'] + [s for _, s in group]
    rows = [[c] + [[(f'{dmean.loc[c, k]:+.2f}', False), (f' ({dpos.loc[c, k]}/3)', True)] for k, _ in group]
            for c in CONDS]
    table(hdr, rows, [2.4] + [(26.0 - 2.4) / len(group)] * len(group), font=9)
    notes(5.5, 26.0)

new_page(False)
heading('Grayscale vs. RGB: average over all 15 test sets',
        'mean Recall@1 over the 15 test sets for each condition and representation')
rows = []
for c in CONDS:
    a, b = mean_rgb.loc[c].mean(), mean.loc[c].mean()
    rows.append([c, f'{a:.2f}', f'{b:.2f}', f'{b - a:+.2f}', str(int((dmean.loc[c] > 0).sum())) + ' of 15'])
table(['Condition', 'RGB', 'Grayscale', 'Gray − RGB', 'Test sets where gray is higher'], rows,
      [3.0, 2.6, 2.6, 2.6, 4.6], font=10)
notes(8, 15.4)

doc.save('results/geo_workbook_gray.docx')
print('saved results/geo_workbook_gray.docx, pages', pno[0] + 1)
